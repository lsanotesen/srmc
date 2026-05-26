from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font
from io import BytesIO
import openpyxl

def parse_excel(file_content):
    wb = load_workbook(filename=BytesIO(file_content), read_only=True)
    ws = wb.active
    
    headers = []
    for col in range(1, ws.max_column + 1):
        header = ws.cell(row=1, column=col).value
        headers.append(header.strip() if header else '')
    
    rows = []
    for row in range(2, ws.max_row + 1):
        row_data = {}
        for col in range(1, ws.max_column + 1):
            header = headers[col - 1]
            value = ws.cell(row=row, column=col).value
            if value is not None:
                row_data[header] = str(value) if isinstance(value, (int, float)) else value
        if row_data:
            rows.append(row_data)
    
    return rows

def generate_excel_template():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "服务导入模板"
    
    headers = [
        'service_name', 'service_code', 'service_type', 'module', 'environment',
        'ip', 'port', 'service_path', 'work_dir', 'start_script',
        'stop_script', 'restart_script', 'log_path', 'check_type',
        'check_keyword', 'pid_file', 'owner', 'remark'
    ]
    
    header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
    header_font = Font(color='FFFFFF', bold=True)
    
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
    
    examples = [
        ['示例应用', 'app-demo', 'APP', 'demo', 'DEV', '192.168.1.10', '8080', 
         '/opt/app/demo', '/opt/app/demo', 'start.sh', 'stop.sh', 'restart.sh', 
         '/var/log/demo', 'PROCESS', 'demo.jar', '/var/run/demo.pid', '张三', '测试应用'],
        ['示例MySQL', 'mysql-prod', 'MYSQL', 'database', 'PROD', '192.168.1.20', '3306',
         '/usr/local/mysql', '/usr/local/mysql', '', '', '',
         '/var/log/mysql', 'PORT', '3306', '', '李四', '生产数据库']
    ]
    
    for row_idx, example in enumerate(examples, 2):
        for col_idx, value in enumerate(example, 1):
            ws.cell(row=row_idx, column=col_idx, value=value)
    
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    
    return output
