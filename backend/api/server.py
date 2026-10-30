import asyncio
import io
import ipaddress
import math
import time
from datetime import datetime
from typing import Optional

from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.datavalidation import DataValidation
from pydantic import BaseModel
from sqlalchemy import false, func
from sqlalchemy.orm import Session

from models.agent import Agent
from models.server import Server, ServerSshAccount
from schemas.server import ServerCreate, ServerUpdate, SshAccountCreate, SshAccountUpdate
from schemas.common import ResponseModel
from services.audit_service import log_audit
from services.data_scope_service import DataScopeService
from services.permission_service import PermissionService
from utils.crypto import encrypt, decrypt
from utils.ssh_pool import ssh_pool
from core.database import get_db
from api.dependencies import require_permission

router = APIRouter()

# 批量导入最大行数
IMPORT_MAX_ROWS = 500

# 手动刷新系统信息：等待 Agent 下一次心跳上报的最长时间（心跳间隔 15s）
REFRESH_TIMEOUT_SECONDS = 20

# 导入模板表头（中文）与字段映射
IMPORT_HEADERS = [
    "主机名", "IP地址", "SSH端口", "SSH用户名", "认证方式",
    "密码", "私钥", "操作系统", "带外类型", "带外地址", "备注",
]
IMPORT_FIELD_MAP = {
    "主机名": "hostname", "hostname": "hostname",
    "IP地址": "ip", "IP": "ip", "ip": "ip",
    "SSH端口": "ssh_port", "ssh_port": "ssh_port",
    "SSH用户名": "username", "username": "username",
    "认证方式": "auth_type", "auth_type": "auth_type",
    "密码": "password", "password": "password",
    "私钥": "private_key", "private_key": "private_key",
    "操作系统": "os_type", "os_type": "os_type",
    "带外类型": "bmc_type", "bmc_type": "bmc_type",
    "带外地址": "bmc_url", "bmc_url": "bmc_url",
    "备注": "remark", "remark": "remark",
}

BMC_TYPES = {"IPMI", "IDRAC", "ILO", "OTHER"}

# SSH 账号导入模板表头与字段映射（与服务器模板一致，兼容中英文表头）
SSH_ACCOUNT_IMPORT_HEADERS = [
    "IP地址", "SSH端口", "SSH用户名", "认证方式", "密码", "私钥", "是否默认", "备注",
]
SSH_ACCOUNT_FIELD_MAP = {
    "IP地址": "ip", "IP": "ip", "ip": "ip",
    "SSH端口": "ssh_port", "ssh_port": "ssh_port",
    "SSH用户名": "username", "username": "username",
    "认证方式": "auth_type", "auth_type": "auth_type",
    "密码": "password", "password": "password",
    "私钥": "private_key", "private_key": "private_key",
    "是否默认": "is_default", "默认账号": "is_default", "is_default": "is_default",
    "备注": "remark", "remark": "remark",
}

# 是否默认列的取值：是/否（兼容 1/0、true/false 等）
SSH_ACCOUNT_TRUE_VALUES = {"是", "1", "true", "yes", "y", "默认"}
SSH_ACCOUNT_FALSE_VALUES = {"", "否", "0", "false", "no", "n", "非默认"}


def _add_auth_type_dropdown(ws, headers: list[str]) -> None:
    """把「认证方式」列做成 Excel 下拉框（密码 / 密钥）。

    模板示例统一使用中文值，避免用户纠结填 key、keyword 还是 password；
    导入解析仍兼容 password/passwd/key/私钥等历史写法。
    """
    column = headers.index("认证方式") + 1
    letter = ws.cell(row=1, column=column).column_letter
    validation = DataValidation(
        type="list",
        formula1='"密码,密钥"',
        allow_blank=True,
        showErrorMessage=True,
        showInputMessage=True,
        errorTitle="认证方式不正确",
        error="请从下拉列表选择「密码」或「密钥」",
        promptTitle="认证方式",
        prompt="密码认证选择「密码」；密钥认证选择「密钥」；留空则按密码/私钥列自动判断",
    )
    ws.add_data_validation(validation)
    validation.add(f"{letter}2:{letter}{IMPORT_MAX_ROWS + 1}")


class BatchIdsPayload(BaseModel):
    ids: list[int]


def _scope_server_query(query, user):
    """组织隔离：超级管理员可见全部，普通用户仅可见本组织服务器。"""
    if DataScopeService._is_super_admin(user):
        return query
    if user.organization_id is not None:
        return query.filter(Server.organization_id == user.organization_id)
    # 无组织的普通用户看不到任何服务器
    return query.filter(false())


def _serialize_server(s: Server, agent_info: Optional[dict] = None, ssh_account_count: int = 0) -> dict:
    """转成可序列化的字典；不含 password / private_key，避免凭据密文外泄。

    agent_info：该 IP 对应 Agent 采集的系统信息（发行版、架构）与安装/在线状态，
    未安装 Agent 时 agent_installed=false（前端回退到服务器登记的 os_type）。
    ssh_account_count：该服务器已配置的 SSH 账号数量，0 表示未配置。
    """
    agent_info = agent_info or {}
    installed = bool(agent_info.get("installed"))
    return {
        "id": s.id,
        "hostname": s.hostname,
        "ip": s.ip,
        "ssh_port": s.ssh_port,
        "username": s.username,
        "os_type": s.os_type,
        "agent_os_distro": agent_info.get("os_distro") or None,
        "agent_arch": agent_info.get("arch") or None,
        # Agent 安装状态：按 IP 关联 agents 表；未安装时在线/心跳均为 null
        "agent_installed": installed,
        "agent_online": bool(agent_info.get("online")) if installed else None,
        "agent_last_heartbeat": agent_info.get("last_heartbeat") if installed else None,
        "bmc_type": s.bmc_type,
        "bmc_url": s.bmc_url,
        "remark": s.remark,
        "has_password": bool(s.password),
        "has_private_key": bool(s.private_key),
        "ssh_account_count": ssh_account_count,
        "organization_id": s.organization_id,
        "created_at": s.created_at,
        "updated_at": s.updated_at,
    }


def _agent_info_map(db: Session, servers: list[Server]) -> dict:
    """按 IP 批量查询 Agent 采集信息（发行版、架构）与安装/在线状态，避免逐台查询。

    同一 IP 存在多条 Agent 记录时，取最近心跳的一条（不看是否有采集数据）。
    """
    ips = {s.ip for s in servers if s.ip}
    if not ips:
        return {}
    # 延迟导入，避免与 api.agent 形成模块级循环依赖
    from api.agent import manager as agent_manager
    rows = db.query(Agent.ip, Agent.os_distro, Agent.arch, Agent.last_heartbeat, Agent.uuid).filter(
        Agent.ip.in_(ips)
    ).all()
    result = {}
    for ip, distro, arch, heartbeat, uuid in rows:
        current = result.get(ip)
        if current is not None:
            # 已有记录：无心跳的不覆盖，有心跳的只保留更新的一条
            if heartbeat is None:
                continue
            cur_hb = current["heartbeat"]
            if cur_hb is not None and heartbeat <= cur_hb:
                continue
        result[ip] = {
            "os_distro": distro,
            "arch": arch,
            "heartbeat": heartbeat,
            "last_heartbeat": heartbeat.isoformat() if heartbeat else None,
            "installed": True,
            "online": agent_manager.is_online(uuid, heartbeat),
        }
    return result


def _account_count_map(db: Session, servers: list[Server]) -> dict:
    """按服务器批量统计 SSH 账号数量，避免列表逐行查询。"""
    ids = [s.id for s in servers]
    if not ids:
        return {}
    rows = db.query(ServerSshAccount.server_id, func.count(ServerSshAccount.id)).filter(
        ServerSshAccount.server_id.in_(ids)
    ).group_by(ServerSshAccount.server_id).all()
    return {server_id: count for server_id, count in rows}


def _default_account(db: Session, server: Server) -> Optional[ServerSshAccount]:
    return db.query(ServerSshAccount).filter(
        ServerSshAccount.server_id == server.id,
        ServerSshAccount.is_default.is_(True),
    ).first()


def _mirror_default_account(server: Server, account: ServerSshAccount) -> None:
    """默认账号变更后同步 servers 原字段，保持 Agent 部署等既有流程可用。"""
    server.username = account.username
    server.password = account.password
    server.private_key = account.private_key


def _set_default_account(db: Session, server: Server, account: ServerSshAccount) -> None:
    """把账号设为默认，并取消同服务器其他账号的默认标记。"""
    db.query(ServerSshAccount).filter(
        ServerSshAccount.server_id == server.id,
        ServerSshAccount.id != account.id,
        ServerSshAccount.is_default.is_(True),
    ).update({ServerSshAccount.is_default: False}, synchronize_session=False)
    account.is_default = True
    _mirror_default_account(server, account)


def _sync_default_account(db: Session, server: Server) -> Optional[ServerSshAccount]:
    """servers 上的凭据被编辑后，同步到默认账号；无账号且有凭据时补建一条。

    该函数保证“服务器原字段 = 默认账号”，供未改造的既有流程继续读取。
    """
    account = _default_account(db, server)
    if not account:
        if not server.password and not server.private_key:
            return None
        account = ServerSshAccount(server_id=server.id, is_default=True)
        db.add(account)
    else:
        duplicated = db.query(ServerSshAccount).filter(
            ServerSshAccount.server_id == server.id,
            ServerSshAccount.username == server.username,
            ServerSshAccount.id != account.id,
        ).first()
        if duplicated:
            raise HTTPException(
                status_code=400,
                detail="该 SSH 用户名已被此服务器的其他账号使用，请先在 SSH 账号管理中调整",
            )
    account.username = server.username
    account.password = server.password
    account.private_key = server.private_key
    return account


def _serialize_ssh_account(account: ServerSshAccount) -> dict:
    """账号对外序列化：只返回是否已配置凭据，不返回密码/私钥。"""
    return {
        "id": account.id,
        "server_id": account.server_id,
        "username": account.username,
        "is_default": bool(account.is_default),
        "has_password": bool(account.password),
        "has_private_key": bool(account.private_key),
        "remark": account.remark,
        "created_at": account.created_at,
        "updated_at": account.updated_at,
    }


def _scope_server_or_404(db: Session, user, server_id: int) -> Server:
    server = _scope_server_query(db.query(Server), user).filter(Server.id == server_id).first()
    if not server:
        raise HTTPException(status_code=404, detail="Server not found")
    return server


def _normalize_account_secrets(password: Optional[str], private_key: Optional[str]) -> tuple[str, str]:
    """账号凭据去空白，密码与私钥互斥。空值表示“不修改”。"""
    password = (password or "").strip()
    private_key = (private_key or "").strip()
    if password and private_key:
        raise HTTPException(status_code=400, detail="密码与私钥不能同时填写")
    return password, private_key


def _normalize_bmc_url(value: Optional[str]) -> Optional[str]:
    """带外地址只允许 http(s) 完整 URL，空值按未配置处理。"""
    if not value:
        return None
    url = value.strip()
    if not url:
        return None
    if not (url.startswith("http://") or url.startswith("https://")):
        raise HTTPException(status_code=400, detail="带外地址必须以 http:// 或 https:// 开头")
    return url


def _normalize_bmc_type(value: Optional[str], has_url: bool) -> Optional[str]:
    if not value:
        return "OTHER" if has_url else None
    bmc_type = value.strip().upper()
    if bmc_type not in BMC_TYPES:
        raise HTTPException(status_code=400, detail="带外类型仅支持 IPMI / iDRAC / iLO / OTHER")
    return bmc_type


def _apply_credentials_filters(data: dict) -> dict:
    """空凭据不落库/不覆盖；填了密码则清除私钥，反之亦然。"""
    password = (data.get("password") or "").strip()
    private_key = (data.get("private_key") or "").strip()
    data.pop("password", None)
    data.pop("private_key", None)
    if password:
        data["password"] = encrypt(password)
        data["private_key"] = None
    elif private_key:
        data["private_key"] = encrypt(private_key)
        data["password"] = None
    return data


@router.get("/servers", response_model=ResponseModel)
async def get_servers(
    db: Session = Depends(get_db),
    user = Depends(require_permission("service_view")),
    page: int = 1,
    page_size: int = 20,
    keyword: str | None = None,
    os_type: str | None = None,
    has_bmc: bool | None = None,
    agent_status: str | None = Query(None, description="Agent 状态筛选：online/offline/no-agent"),
):
    page = max(page, 1)
    page_size = min(max(page_size, 1), 100)
    agent_status = (agent_status or "").strip().lower() or None
    if agent_status and agent_status not in ("online", "offline", "no-agent"):
        raise HTTPException(status_code=400, detail="agent_status 仅支持 online/offline/no-agent")

    query = _scope_server_query(db.query(Server), user)
    # total_all：数据范围内的服务器总数（不受筛选条件影响），供页面区分"总数 / 筛选结果"
    total_all = query.count()
    if keyword:
        like = f"%{keyword.strip()}%"
        query = query.filter(
            (Server.hostname.ilike(like))
            | (Server.ip.ilike(like))
            | (Server.username.ilike(like))
            | (Server.remark.ilike(like))
        )
    if os_type:
        query = query.filter(Server.os_type == os_type.upper())
    if has_bmc is True:
        query = query.filter(Server.bmc_url.isnot(None), Server.bmc_url != "")
    elif has_bmc is False:
        query = query.filter((Server.bmc_url.is_(None)) | (Server.bmc_url == ""))

    # 按 IP 数值升序，避免字符串排序导致 .10 排在 .2 前面
    ordered_query = query.order_by(func.inet6_aton(Server.ip), Server.ip)

    if agent_status:
        # Agent 状态不在 DB 中，只能后置过滤：先取出 DB 过滤范围内的全部服务器，
        # 再按 Agent 状态筛选，最后计算 total 与分页
        candidates = ordered_query.all()
        all_info = _agent_info_map(db, candidates)

        def _match(s: Server) -> bool:
            info = all_info.get(s.ip) or {}
            if not info.get("installed"):
                return agent_status == "no-agent"
            if agent_status == "no-agent":
                return False
            return bool(info.get("online")) if agent_status == "online" else not bool(info.get("online"))

        filtered = [s for s in candidates if _match(s)]
        total = len(filtered)
        servers = filtered[(page - 1) * page_size: page * page_size]
        info_map = {s.ip: all_info[s.ip] for s in servers if s.ip in all_info}
    else:
        total = query.count()
        servers = ordered_query.offset((page - 1) * page_size).limit(page_size).all()
        info_map = _agent_info_map(db, servers)

    count_map = _account_count_map(db, servers)
    return ResponseModel(data={
        "items": [
            _serialize_server(s, info_map.get(s.ip), count_map.get(s.id, 0))
            for s in servers
        ],
        "total": total, "total_all": total_all, "page": page, "page_size": page_size,
    })


@router.get("/servers/{server_id}", response_model=ResponseModel)
async def get_server(server_id: int, db: Session = Depends(get_db), user = Depends(require_permission("service_view"))):
    server = _scope_server_query(db.query(Server), user).filter(Server.id == server_id).first()
    if not server:
        raise HTTPException(status_code=404, detail="Server not found")
    info_map = _agent_info_map(db, [server])
    count_map = _account_count_map(db, [server])
    return ResponseModel(data=_serialize_server(server, info_map.get(server.ip), count_map.get(server.id, 0)))


@router.post("/servers", response_model=ResponseModel)
async def create_server(server_create: ServerCreate, db: Session = Depends(get_db), user = Depends(require_permission("host:manage"))):
    existing = db.query(Server).filter(
        Server.ip == str(server_create.ip),
        Server.organization_id == user.organization_id,
    ).first()
    if existing:
        raise HTTPException(status_code=400, detail="该 IP 在当前组织已存在")

    bmc_url = _normalize_bmc_url(server_create.bmc_url)
    bmc_type = _normalize_bmc_type(server_create.bmc_type, bool(bmc_url))

    server = Server(
        hostname=server_create.hostname.strip(),
        ip=str(server_create.ip),
        ssh_port=server_create.ssh_port,
        username=server_create.username.strip(),
        password=encrypt(server_create.password) if server_create.password else None,
        private_key=encrypt(server_create.private_key) if server_create.private_key else None,
        os_type=server_create.os_type,
        bmc_type=bmc_type,
        bmc_url=bmc_url,
        remark=(server_create.remark or "").strip() or None,
        organization_id=user.organization_id
    )
    db.add(server)
    db.flush()
    # 新建服务器时把凭据落为默认账号，保证账号列与既有流程一致
    _sync_default_account(db, server)
    db.commit()

    log_audit(db, user.id, user.username, "CREATE", server_id=server.id, ip=server.ip, result="success")

    return ResponseModel(data=_serialize_server(server))


@router.put("/servers/{server_id}", response_model=ResponseModel)
async def update_server(server_id: int, server_update: ServerUpdate, db: Session = Depends(get_db), user = Depends(require_permission("host:manage"))):
    server = _scope_server_query(db.query(Server), user).filter(Server.id == server_id).first()
    if not server:
        raise HTTPException(status_code=404, detail="Server not found")

    data = server_update.dict(exclude_unset=True)

    # IP 变更时校验组织内唯一
    if "ip" in data and data["ip"] is not None:
        new_ip = str(data["ip"])
        duplicated = db.query(Server).filter(
            Server.ip == new_ip,
            Server.organization_id == server.organization_id,
            Server.id != server.id,
        ).first()
        if duplicated:
            raise HTTPException(status_code=400, detail="该 IP 在当前组织已存在")
        data["ip"] = new_ip

    if "bmc_url" in data:
        data["bmc_url"] = _normalize_bmc_url(data["bmc_url"])
    if "bmc_type" in data or "bmc_url" in data:
        final_url = data.get("bmc_url", server.bmc_url)
        final_type = data.get("bmc_type", server.bmc_type)
        data["bmc_type"] = _normalize_bmc_type(final_type, bool(final_url))

    if "remark" in data:
        data["remark"] = (data.get("remark") or "").strip() or None

    # 关键：空密码/私钥不覆盖已有凭据（避免编辑时把密码清空）
    data = _apply_credentials_filters(data)

    for key, value in data.items():
        setattr(server, key, value)

    # 编辑服务器维护的是默认账号，需同步到账号表
    _sync_default_account(db, server)
    db.commit()

    log_audit(db, user.id, user.username, "UPDATE", server_id=server.id, ip=server.ip, result="success")

    return ResponseModel(data=_serialize_server(server))


@router.delete("/servers/{server_id}", response_model=ResponseModel)
async def delete_server(server_id: int, db: Session = Depends(get_db), user = Depends(require_permission("host:manage"))):
    server = _scope_server_query(db.query(Server), user).filter(Server.id == server_id).first()
    if not server:
        raise HTTPException(status_code=404, detail="Server not found")

    ip = server.ip
    db.delete(server)
    db.commit()

    log_audit(db, user.id, user.username, "DELETE", server_id=server_id, ip=ip, result="success")

    return ResponseModel(message="Server deleted successfully")


@router.post("/servers/batch-delete", response_model=ResponseModel)
async def batch_delete_servers(
    payload: BatchIdsPayload,
    db: Session = Depends(get_db),
    user = Depends(require_permission("host:manage")),
):
    """批量删除服务器：逐条做组织隔离校验，返回成功与失败明细。"""
    ids = list(dict.fromkeys(payload.ids or []))
    if not ids:
        raise HTTPException(status_code=400, detail="请选择要删除的服务器")

    servers = _scope_server_query(db.query(Server), user).filter(Server.id.in_(ids)).all()
    found = {s.id: s for s in servers}
    success_ids, failed = [], []
    for server_id in ids:
        server = found.get(server_id)
        if not server:
            failed.append({"id": server_id, "message": "服务器不存在或无权操作"})
            continue
        info = {"id": server.id, "hostname": server.hostname, "ip": server.ip}
        db.delete(server)
        success_ids.append(info)

    if success_ids:
        db.commit()
        for item in success_ids:
            log_audit(db, user.id, user.username, "DELETE", server_id=item["id"], ip=item["ip"], result="success")

    return ResponseModel(data={
        "success": success_ids,
        "failed": failed,
        "success_count": len(success_ids),
        "failed_count": len(failed),
    })


@router.post("/servers/{server_id}/test-connection", response_model=ResponseModel)
async def test_connection(server_id: int, db: Session = Depends(get_db), user = Depends(require_permission("service_view"))):
    server = _scope_server_query(db.query(Server), user).filter(Server.id == server_id).first()
    if not server:
        raise HTTPException(status_code=404, detail="Server not found")

    if not server.password and not server.private_key:
        return ResponseModel(code=1, message="该服务器未配置密码或私钥")

    password = decrypt(server.password) if server.password else None
    private_key = decrypt(server.private_key) if server.private_key else None

    try:
        conn = await ssh_pool.get_connection(server.ip, server.ssh_port, server.username, password, private_key)
        if conn:
            output, error, success = await conn.execute_command("echo 'Connection test successful'")
            if success:
                log_audit(db, user.id, user.username, "LOGIN", server_id=server.id, ip=server.ip, result="success")
                return ResponseModel(message="连接成功")
            log_audit(db, user.id, user.username, "LOGIN", server_id=server.id, ip=server.ip,
                      result="failed", output=(error or output or "")[:500])
            return ResponseModel(code=1, message=f"连接失败：{(error or '命令执行失败')[:200]}")
    except Exception as e:
        log_audit(db, user.id, user.username, "LOGIN", server_id=server.id, ip=server.ip,
                  result="failed", output=str(e)[:500])
        return ResponseModel(code=1, message=f"连接失败：{str(e)[:200]}")

    return ResponseModel(code=1, message="连接失败：无法建立 SSH 连接")


@router.post("/servers/refresh-system-info", response_model=ResponseModel)
async def refresh_server_system_info(
    payload: BatchIdsPayload,
    db: Session = Depends(get_db),
    user = Depends(require_permission("host:manage")),
):
    """手动刷新选中服务器的操作系统信息。

    不通过 SSH 探测：复用 Agent 心跳采集，等待 Agent 下一次心跳上报后返回真实结果。
    Agent 未安装 / 离线 / 超时均明确反馈；上报未完成前不会返回成功。
    """
    from api.agent import manager as agent_manager

    ids = list(dict.fromkeys(payload.ids or []))
    if not ids:
        raise HTTPException(status_code=400, detail="请选择要刷新系统信息的服务器")

    servers = _scope_server_query(db.query(Server), user).filter(Server.id.in_(ids)).all()
    found = {s.id: s for s in servers}
    result_by_id = {}
    pending = []

    for server_id in ids:
        server = found.get(server_id)
        if not server:
            result_by_id[server_id] = {
                "id": server_id, "hostname": "", "ip": "", "status": "failed",
                "os_distro": None, "message": "服务器不存在或无权操作",
            }
            continue
        # 同一 IP 可能有多条 Agent 记录，优先取最近心跳的一条
        agent = db.query(Agent).filter(Agent.ip == server.ip).order_by(Agent.last_heartbeat.desc()).first()
        if not agent:
            result_by_id[server_id] = {
                "id": server.id, "hostname": server.hostname, "ip": server.ip,
                "status": "no_agent", "os_distro": None, "message": "该服务器未安装 Agent",
            }
        elif not agent_manager.is_online(agent.uuid, agent.last_heartbeat):
            result_by_id[server_id] = {
                "id": server.id, "hostname": server.hostname, "ip": server.ip,
                "status": "offline", "os_distro": agent.os_distro or None,
                "message": "Agent 离线，无法采集",
            }
        else:
            pending.append({
                "id": server.id, "hostname": server.hostname, "ip": server.ip, "uuid": agent.uuid,
            })

    if pending:
        # 心跳间隔 15 秒，留出重连与调度余量
        baseline = datetime.now()
        uuids = [item["uuid"] for item in pending]
        latest = {}
        deadline = time.time() + REFRESH_TIMEOUT_SECONDS
        while time.time() < deadline:
            # MySQL 默认 REPEATABLE READ，需结束当前只读事务才能看到心跳线程提交的新值
            db.rollback()
            rows = db.query(Agent.uuid, Agent.last_heartbeat, Agent.os_distro).filter(Agent.uuid.in_(uuids)).all()
            latest = {uuid: (heartbeat, distro) for uuid, heartbeat, distro in rows}
            if all(latest.get(uuid) and latest[uuid][0] and latest[uuid][0] > baseline for uuid in uuids):
                break
            await asyncio.sleep(1)

        db.rollback()
        rows = db.query(Agent.uuid, Agent.last_heartbeat, Agent.os_distro).filter(Agent.uuid.in_(uuids)).all()
        latest = {uuid: (heartbeat, distro) for uuid, heartbeat, distro in rows}

        for item in pending:
            heartbeat, distro = latest.get(item["uuid"], (None, None))
            reported = bool(heartbeat and heartbeat > baseline)
            if reported and distro:
                status, message = "success", "已获取最新操作系统信息"
            elif reported:
                status, message = "failed", "Agent 已上报，但未采集到发行版信息"
            else:
                status, message = "failed", f"等待 Agent 上报超时（{REFRESH_TIMEOUT_SECONDS} 秒）"
            result_by_id[item["id"]] = {
                "id": item["id"], "hostname": item["hostname"], "ip": item["ip"],
                "status": status, "os_distro": distro or None, "message": message,
            }

    items = [result_by_id[i] for i in ids]
    success_count = sum(1 for item in items if item["status"] == "success")
    offline_count = sum(1 for item in items if item["status"] == "offline")
    log_audit(
        db, user.id, user.username, "AGENT_CONTROL",
        result="success" if success_count == len(items) else "partial",
        output=f"刷新服务器系统信息 总数={len(items)} 成功={success_count} 离线={offline_count}",
    )

    return ResponseModel(data={
        "total": len(items),
        "success_count": success_count,
        "offline_count": offline_count,
        "failed_count": len(items) - success_count - offline_count,
        "items": items,
    })


def _cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float):
        # Excel 空单元格/异常数据可能读成 NaN；整数值浮点去掉 .0，避免 "22.0" 这类脏数据
        if math.isnan(value) or math.isinf(value):
            return ""
        if value.is_integer():
            value = int(value)
    return str(value).strip()


# 密码/私钥单元格中不能当作真实凭据的占位值
SECRET_PLACEHOLDERS = {"password", "changeme", "登录密码", "示例密码"}
MASK_CHARS = set("*•·●○")


def _secret_text(value) -> str:
    """读取密码/私钥单元格：空值、NaN、纯掩码字符、模板占位词一律视为空。"""
    text = _cell_text(value)
    if not text:
        return ""
    if set(text) <= MASK_CHARS:
        return ""
    if text.lower() in SECRET_PLACEHOLDERS:
        return ""
    return text


def _parse_import_row(row: dict) -> tuple[dict, str]:
    """校验单行导入数据，返回 (数据, 错误信息)。"""
    hostname = _cell_text(row.get("hostname"))
    ip_raw = _cell_text(row.get("ip"))
    username = _cell_text(row.get("username"))

    if not hostname:
        return {}, "主机名不能为空"
    if len(hostname) > 128:
        return {}, "主机名不能超过 128 个字符"
    if not ip_raw:
        return {}, "IP 地址不能为空"
    try:
        ip = str(ipaddress.ip_address(ip_raw))
    except ValueError:
        return {}, f"IP 地址不合法：{ip_raw}"
    if not username:
        return {}, "SSH 用户名不能为空"
    if len(username) > 64:
        return {}, "SSH 用户名不能超过 64 个字符"

    port_raw = _cell_text(row.get("ssh_port")) or "22"
    try:
        port = int(float(port_raw))
    except (TypeError, ValueError):
        return {}, f"SSH 端口不合法：{port_raw}"
    if not (1 <= port <= 65535):
        return {}, "SSH 端口需在 1-65535 之间"

    auth_raw = _cell_text(row.get("auth_type")).lower()
    password = _secret_text(row.get("password"))
    private_key = _secret_text(row.get("private_key"))
    if auth_raw in ("password", "passwd", "密码"):
        auth_type = "password"
    elif auth_raw in ("key", "密钥", "私钥"):
        auth_type = "key"
    else:
        auth_type = "password" if password else ("key" if private_key else "")
    if auth_type == "password" and not password:
        return {}, "认证方式为密码时必须填写密码"
    if auth_type == "key" and not private_key:
        return {}, "认证方式为密钥时必须填写私钥"

    os_raw = _cell_text(row.get("os_type")) or "LINUX"
    os_type = os_raw.upper()
    if os_type not in ("LINUX", "UNIX"):
        return {}, f"操作系统仅支持 LINUX / UNIX，当前：{os_raw}"

    bmc_url_raw = _cell_text(row.get("bmc_url"))
    if bmc_url_raw and not (bmc_url_raw.startswith("http://") or bmc_url_raw.startswith("https://")):
        return {}, "带外地址必须以 http:// 或 https:// 开头"
    bmc_type_raw = _cell_text(row.get("bmc_type")).upper()
    if bmc_type_raw and bmc_type_raw not in BMC_TYPES:
        return {}, "带外类型仅支持 IPMI / iDRAC / iLO / OTHER"
    bmc_type = bmc_type_raw or ("OTHER" if bmc_url_raw else None)

    remark = _cell_text(row.get("remark"))
    if len(remark) > 255:
        return {}, "备注不能超过 255 个字符"

    return {
        "hostname": hostname,
        "ip": ip,
        "ssh_port": port,
        "username": username,
        "password": password or None,
        "private_key": private_key or None,
        "os_type": os_type,
        "bmc_type": bmc_type,
        "bmc_url": bmc_url_raw or None,
        "remark": remark or None,
    }, ""


def _secret_matches(stored: Optional[str], plain: str) -> bool:
    """比较库内凭据密文与 Excel 明文是否一致（解密失败按不一致处理）。"""
    if not stored:
        return False
    try:
        return (decrypt(stored) or "").strip() == plain.strip()
    except Exception:
        return False


def _build_import_update(server: Server, row: dict) -> tuple[dict, list[str], str]:
    """按 IP 命中已有服务器时，构建更新字段。

    只收集"Excel 单元格有值且与数据库不同"的字段；空单元格一律保留原值。
    凭据与库内解密值比对，不一致才更新，保证同一文件重复导入幂等。
    返回 (待更新字段, 变更字段中文名, 错误信息)。
    """
    changes: dict = {}
    changed_fields: list[str] = []

    hostname = _cell_text(row.get("hostname"))
    if hostname:
        if len(hostname) > 128:
            return {}, [], "主机名不能超过 128 个字符"
        if hostname != server.hostname:
            changes["hostname"] = hostname
            changed_fields.append("主机名")

    username = _cell_text(row.get("username"))
    if username:
        if len(username) > 64:
            return {}, [], "SSH 用户名不能超过 64 个字符"
        if username != server.username:
            changes["username"] = username
            changed_fields.append("SSH 用户名")

    port_raw = _cell_text(row.get("ssh_port"))
    if port_raw:
        try:
            port = int(float(port_raw))
        except (TypeError, ValueError):
            return {}, [], f"SSH 端口不合法：{port_raw}"
        if not (1 <= port <= 65535):
            return {}, [], "SSH 端口需在 1-65535 之间"
        if port != server.ssh_port:
            changes["ssh_port"] = port
            changed_fields.append("SSH 端口")

    os_raw = _cell_text(row.get("os_type"))
    if os_raw:
        os_type = os_raw.upper()
        if os_type not in ("LINUX", "UNIX"):
            return {}, [], f"操作系统仅支持 LINUX / UNIX，当前：{os_raw}"
        if os_type != server.os_type:
            changes["os_type"] = os_type
            changed_fields.append("操作系统")

    bmc_url_raw = _cell_text(row.get("bmc_url"))
    bmc_type_raw = _cell_text(row.get("bmc_type"))
    if bmc_url_raw and not (bmc_url_raw.startswith("http://") or bmc_url_raw.startswith("https://")):
        return {}, [], "带外地址必须以 http:// 或 https:// 开头"
    bmc_type = None
    if bmc_type_raw:
        bmc_type = bmc_type_raw.upper()
        if bmc_type not in BMC_TYPES:
            return {}, [], "带外类型仅支持 IPMI / iDRAC / iLO / OTHER"
    # 类型未填但填了地址、且原来没有类型时，按现有规则补 OTHER
    final_bmc_type = bmc_type or server.bmc_type or ("OTHER" if bmc_url_raw else None)
    if bmc_url_raw and bmc_url_raw != (server.bmc_url or ""):
        changes["bmc_url"] = bmc_url_raw
        changed_fields.append("带外地址")
    if final_bmc_type != server.bmc_type and (bmc_type_raw or bmc_url_raw):
        changes["bmc_type"] = final_bmc_type
        changed_fields.append("带外类型")

    remark = _cell_text(row.get("remark"))
    if remark:
        if len(remark) > 255:
            return {}, [], "备注不能超过 255 个字符"
        if remark != (server.remark or ""):
            changes["remark"] = remark
            changed_fields.append("备注")

    # 凭据：Excel 单元格有真实值且与库内现值不一致时才更新（支持通过 Excel 改密码），
    # 空值/掩码/占位词不落库；填了密码则清空私钥，反之亦然（与编辑接口一致）
    auth_raw = _cell_text(row.get("auth_type")).lower()
    if auth_raw and auth_raw not in ("password", "passwd", "密码", "key", "密钥", "私钥"):
        return {}, [], "认证方式仅支持 密码 / 密钥"
    password = _secret_text(row.get("password"))
    private_key = _secret_text(row.get("private_key"))
    if password and private_key:
        return {}, [], "密码与私钥不能同时填写"
    if password:
        if not _secret_matches(server.password, password):
            changes["password"] = encrypt(password)
            changes["private_key"] = None
            changed_fields.append("密码")
    elif private_key:
        if not _secret_matches(server.private_key, private_key):
            changes["private_key"] = encrypt(private_key)
            changes["password"] = None
            changed_fields.append("私钥")

    return changes, changed_fields, ""


@router.get("/servers/import/template")
async def download_server_import_template(_user=Depends(require_permission("host:manage"))):
    """下载服务器批量导入模板（xlsx）。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "服务器导入"
    ws.append(IMPORT_HEADERS)
    ws.append([
        "web-node-01", "192.168.1.10", 22, "root", "密码",
        "登录密码", "", "LINUX", "iDRAC", "https://192.168.1.110", "示例行，导入前请删除",
    ])
    widths = [16, 16, 10, 14, 12, 16, 24, 12, 12, 30, 26]
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=index).column_letter].width = width
    # 认证方式列提供下拉选项，避免手填 key/keyword/passwd 等歧义值
    _add_auth_type_dropdown(ws, IMPORT_HEADERS)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=server-import-template.xlsx"},
    )


@router.post("/servers/batch-import", response_model=ResponseModel)
async def batch_import_servers(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    user = Depends(require_permission("host:manage")),
):
    """批量导入服务器（xlsx）。

    按 IP 匹配：已存在（当前组织内）则更新有值且发生变化的字段，不存在则新增；
    空单元格保留数据库原值。每行独立提交，单行失败不影响其他合法记录。
    """
    file_content = await file.read()
    if not file_content:
        raise HTTPException(status_code=400, detail="请上传文件")
    try:
        # data_only=True：公式单元格取 Excel 缓存的计算结果，避免把 "=A2" 这类公式文本当成密码/私钥落库
        wb = load_workbook(filename=io.BytesIO(file_content), read_only=True, data_only=True)
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"文件解析失败: {e}")

    ws = wb.active
    header_cells = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
    field_by_index = {}
    for index, cell in enumerate(header_cells):
        field = IMPORT_FIELD_MAP.get(_cell_text(cell))
        if field and field not in field_by_index.values():
            field_by_index[index] = field
    if "hostname" not in field_by_index.values() or "ip" not in field_by_index.values():
        raise HTTPException(status_code=400, detail="模板表头不正确，请下载最新模板")

    # 当前数据范围内按 IP 建立索引，用于判断"更新"还是"新增"
    existing_by_ip = {}
    for server in _scope_server_query(db.query(Server), user).all():
        existing_by_ip.setdefault(server.ip, server)

    seen_ips = set()
    items = []
    added = updated = unchanged = failed = 0

    for row_index, values in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if row_index - 1 > IMPORT_MAX_ROWS:
            raise HTTPException(status_code=400, detail=f"单次最多导入 {IMPORT_MAX_ROWS} 行")
        if not any(_cell_text(v) for v in values):
            continue

        row = {field: values[index] if index < len(values) else None
               for index, field in field_by_index.items()}
        item = {
            "row": row_index,
            "hostname": _cell_text(row.get("hostname")),
            "ip": _cell_text(row.get("ip")),
        }
        if not item["ip"]:
            failed += 1
            items.append({**item, "status": "failed", "message": "IP 地址不能为空"})
            continue
        try:
            ip = str(ipaddress.ip_address(item["ip"]))
        except ValueError:
            failed += 1
            items.append({**item, "status": "failed", "message": f"IP 地址不合法：{item['ip']}"})
            continue
        item["ip"] = ip

        if ip in seen_ips:
            failed += 1
            items.append({**item, "status": "failed", "message": "文件内 IP 重复"})
            continue
        seen_ips.add(ip)

        server = existing_by_ip.get(ip)
        if server:
            # 已存在：按 IP 更新有值且发生变化的字段
            changes, changed_fields, error = _build_import_update(server, row)
            if error:
                failed += 1
                items.append({**item, "status": "failed", "message": error})
                continue
            if not changes:
                unchanged += 1
                items.append({
                    **item, "hostname": server.hostname,
                    "status": "unchanged", "message": "无变化",
                })
                continue
            try:
                for key, value in changes.items():
                    setattr(server, key, value)
                # 服务器凭据更新后同步默认账号
                _sync_default_account(db, server)
                db.commit()
                updated += 1
                items.append({
                    **item, "hostname": server.hostname,
                    "status": "updated", "message": "已更新：" + "、".join(changed_fields),
                })
            except Exception as e:
                db.rollback()
                failed += 1
                items.append({
                    **item, "hostname": server.hostname,
                    "status": "failed", "message": f"更新失败：{e}",
                })
            continue

        # 不存在：按新增行完整校验后写入
        data, error = _parse_import_row(row)
        if error:
            failed += 1
            items.append({**item, "status": "failed", "message": error})
            continue
        try:
            new_server = Server(
                hostname=data["hostname"],
                ip=data["ip"],
                ssh_port=data["ssh_port"],
                username=data["username"],
                password=encrypt(data["password"]) if data["password"] else None,
                private_key=encrypt(data["private_key"]) if data["private_key"] else None,
                os_type=data["os_type"],
                bmc_type=data["bmc_type"],
                bmc_url=data["bmc_url"],
                remark=data["remark"],
                organization_id=user.organization_id,
            )
            db.add(new_server)
            db.flush()
            _sync_default_account(db, new_server)
            db.commit()
            added += 1
            items.append({
                **item, "hostname": new_server.hostname,
                "status": "added", "message": "新增成功",
            })
        except Exception as e:
            db.rollback()
            failed += 1
            items.append({**item, "status": "failed", "message": f"新增失败：{e}"})

    log_audit(
        db, user.id, user.username, "IMPORT",
        ip=None, result="success" if failed == 0 else "partial",
        output=f"file={file.filename}, added={added}, updated={updated}, "
               f"unchanged={unchanged}, failed={failed}",
    )

    return ResponseModel(data={
        "total": added + updated + unchanged + failed,
        "added": added,
        "updated": updated,
        "unchanged": unchanged,
        "failed": failed,
        "items": items,
    })


# ========================= SSH 账号管理 =========================


@router.get("/servers/{server_id}/usable-ssh-accounts", response_model=ResponseModel)
async def get_usable_ssh_accounts(
    server_id: int,
    db: Session = Depends(get_db),
    user = Depends(require_permission("service_view")),
):
    """获取当前用户有权使用的 SSH 账号，供列表「登录」选择。

    权限规则沿用现有体系：需 webshell:login；用户名为 root 的账号
    额外需要 webshell:root-login；未配置密码/私钥的账号不可登录。
    真正的登录校验仍在 WebShell 建链时执行，前端过滤不作为安全边界。
    """
    server = _scope_server_or_404(db, user, server_id)
    accounts = db.query(ServerSshAccount).filter(
        ServerSshAccount.server_id == server.id
    ).order_by(ServerSshAccount.is_default.desc(), ServerSshAccount.id).all()

    server_info = {
        "id": server.id,
        "hostname": server.hostname,
        "ip": server.ip,
        "ssh_port": server.ssh_port,
    }

    if not PermissionService.user_has_permission(db, user.id, "webshell:login"):
        return ResponseModel(data={
            "server": server_info, "items": [], "total": 0,
            "message": "当前用户无 WebShell 登录权限",
        })

    can_root = PermissionService.user_has_permission(db, user.id, "webshell:root-login")
    with_credential = [a for a in accounts if a.password or a.private_key]
    usable = [
        a for a in with_credential
        if (a.username or "").strip() != "root" or can_root
    ]

    message = ""
    if not usable:
        if not accounts:
            message = "该服务器尚未配置 SSH 账号，请先在 SSH 账号管理中配置"
        elif not with_credential:
            message = "已配置的 SSH 账号均未设置密码或私钥，无法登录"
        else:
            message = "可用账号均为 root，当前用户缺少 Root 登录权限，请联系管理员授权"

    return ResponseModel(data={
        "server": server_info,
        "items": [_serialize_ssh_account(a) for a in usable],
        "total": len(usable),
        "root_blocked": len(with_credential) - len(usable),
        "message": message,
    })


@router.get("/servers/{server_id}/ssh-accounts", response_model=ResponseModel)
async def get_ssh_accounts(
    server_id: int,
    db: Session = Depends(get_db),
    user = Depends(require_permission("service_view")),
):
    """查看服务器 SSH 账号列表（不返回密码/私钥）。"""
    server = _scope_server_or_404(db, user, server_id)
    accounts = db.query(ServerSshAccount).filter(
        ServerSshAccount.server_id == server.id
    ).order_by(ServerSshAccount.is_default.desc(), ServerSshAccount.id).all()
    return ResponseModel(data={
        "server": {
            "id": server.id,
            "hostname": server.hostname,
            "ip": server.ip,
            "ssh_port": server.ssh_port,
        },
        "items": [_serialize_ssh_account(a) for a in accounts],
        "total": len(accounts),
    })


@router.post("/servers/{server_id}/ssh-accounts", response_model=ResponseModel)
async def create_ssh_account(
    server_id: int,
    payload: SshAccountCreate,
    db: Session = Depends(get_db),
    user = Depends(require_permission("host:manage")),
):
    """新增 SSH 账号：首个账号自动成为默认账号。"""
    server = _scope_server_or_404(db, user, server_id)

    username = (payload.username or "").strip()
    if not username:
        raise HTTPException(status_code=400, detail="SSH 用户名不能为空")
    password, private_key = _normalize_account_secrets(payload.password, payload.private_key)
    if not password and not private_key:
        raise HTTPException(status_code=400, detail="新增账号请填写密码或私钥")

    duplicated = db.query(ServerSshAccount).filter(
        ServerSshAccount.server_id == server.id,
        ServerSshAccount.username == username,
    ).first()
    if duplicated:
        raise HTTPException(status_code=400, detail="该服务器已存在同名 SSH 账号")

    account_count = db.query(ServerSshAccount).filter(
        ServerSshAccount.server_id == server.id
    ).count()

    account = ServerSshAccount(
        server_id=server.id,
        username=username,
        password=encrypt(password) if password else None,
        private_key=encrypt(private_key) if private_key else None,
        is_default=False,
        remark=(payload.remark or "").strip() or None,
    )
    db.add(account)
    db.flush()

    # 首个账号或显式勾选默认时设为默认，并回写 servers 兼容字段
    if payload.is_default or account_count == 0:
        _set_default_account(db, server, account)

    db.commit()
    log_audit(
        db, user.id, user.username, "CREATE", server_id=server.id, ip=server.ip,
        result="success", output=f"SSH账号: {username}",
    )
    return ResponseModel(data=_serialize_ssh_account(account))


@router.put("/servers/{server_id}/ssh-accounts/{account_id}", response_model=ResponseModel)
async def update_ssh_account(
    server_id: int,
    account_id: int,
    payload: SshAccountUpdate,
    db: Session = Depends(get_db),
    user = Depends(require_permission("host:manage")),
):
    """更新 SSH 账号：密码/私钥留空表示保留原值。"""
    server = _scope_server_or_404(db, user, server_id)
    account = db.query(ServerSshAccount).filter(
        ServerSshAccount.id == account_id,
        ServerSshAccount.server_id == server.id,
    ).first()
    if not account:
        raise HTTPException(status_code=404, detail="SSH 账号不存在")

    data = payload.dict(exclude_unset=True)

    if "username" in data:
        username = (data.get("username") or "").strip()
        if not username:
            raise HTTPException(status_code=400, detail="SSH 用户名不能为空")
        if username != account.username:
            duplicated = db.query(ServerSshAccount).filter(
                ServerSshAccount.server_id == server.id,
                ServerSshAccount.username == username,
                ServerSshAccount.id != account.id,
            ).first()
            if duplicated:
                raise HTTPException(status_code=400, detail="该服务器已存在同名 SSH 账号")
            account.username = username

    password, private_key = _normalize_account_secrets(data.get("password"), data.get("private_key"))
    if password:
        account.password = encrypt(password)
        account.private_key = None
    elif private_key:
        account.private_key = encrypt(private_key)
        account.password = None

    if "remark" in data:
        account.remark = (data.get("remark") or "").strip() or None

    if data.get("is_default") is True:
        _set_default_account(db, server, account)
    elif data.get("is_default") is False and account.is_default:
        raise HTTPException(
            status_code=400,
            detail="服务器必须保留一个默认账号，请先将其他账号设为默认",
        )
    elif account.is_default:
        # 默认账号的用户名或凭据变更后，同步 servers 兼容字段
        _mirror_default_account(server, account)

    db.commit()
    log_audit(
        db, user.id, user.username, "UPDATE", server_id=server.id, ip=server.ip,
        result="success", output=f"SSH账号: {account.username}",
    )
    return ResponseModel(data=_serialize_ssh_account(account))


@router.delete("/servers/{server_id}/ssh-accounts/{account_id}", response_model=ResponseModel)
async def delete_ssh_account(
    server_id: int,
    account_id: int,
    db: Session = Depends(get_db),
    user = Depends(require_permission("host:manage")),
):
    """删除 SSH 账号；删除默认账号时自动提升最早创建的其他账号。"""
    server = _scope_server_or_404(db, user, server_id)
    account = db.query(ServerSshAccount).filter(
        ServerSshAccount.id == account_id,
        ServerSshAccount.server_id == server.id,
    ).first()
    if not account:
        raise HTTPException(status_code=404, detail="SSH 账号不存在")

    username = account.username
    was_default = bool(account.is_default)
    db.delete(account)
    db.flush()

    if was_default:
        next_account = db.query(ServerSshAccount).filter(
            ServerSshAccount.server_id == server.id
        ).order_by(ServerSshAccount.id).first()
        if next_account:
            _set_default_account(db, server, next_account)
        else:
            # 已无任何账号：清空兼容字段中的凭据，页面显示“未配置”
            server.password = None
            server.private_key = None

    db.commit()
    log_audit(
        db, user.id, user.username, "DELETE", server_id=server.id, ip=server.ip,
        result="success", output=f"SSH账号: {username}",
    )
    return ResponseModel(message="SSH 账号已删除")


@router.get("/servers/ssh-accounts/import/template")
async def download_ssh_account_import_template(_user=Depends(require_permission("host:manage"))):
    """下载 SSH 账号批量导入模板（xlsx）。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "SSH账号导入"
    ws.append(SSH_ACCOUNT_IMPORT_HEADERS)
    # 示例：同一 IP 可重复出现，每行一个账号；SSH 端口需与服务器登记一致
    ws.append(["192.168.1.10", 22, "ops", "密码", "登录密码", "", "是", "示例行，导入前请删除"])
    ws.append(["192.168.1.10", 22, "dba", "密钥", "", "私钥内容", "否", "示例行，导入前请删除"])
    widths = [16, 10, 16, 12, 16, 24, 12, 26]
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=index).column_letter].width = width
    # 认证方式列提供下拉选项（密码 / 密钥）
    _add_auth_type_dropdown(ws, SSH_ACCOUNT_IMPORT_HEADERS)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=ssh-account-import-template.xlsx"},
    )


@router.post("/servers/ssh-accounts/batch-import", response_model=ResponseModel)
async def batch_import_ssh_accounts(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    user = Depends(require_permission("host:manage")),
):
    """批量导入 SSH 账号（xlsx）。

    先对全量行做校验，任一行不合法则整批拒绝、不写入数据；
    校验通过后按 IP + 用户名新增或更新账号，密码/私钥留空保留原值。
    """
    file_content = await file.read()
    if not file_content:
        raise HTTPException(status_code=400, detail="请上传文件")
    try:
        # data_only=True：避免公式单元格被当作密码/私钥文本导入
        wb = load_workbook(filename=io.BytesIO(file_content), read_only=True, data_only=True)
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"文件解析失败: {e}")

    ws = wb.active
    header_cells = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
    field_by_index = {}
    for index, cell in enumerate(header_cells):
        field = SSH_ACCOUNT_FIELD_MAP.get(_cell_text(cell))
        if field and field not in field_by_index.values():
            field_by_index[index] = field
    mapped = set(field_by_index.values())
    if "ip" not in mapped or "username" not in mapped:
        raise HTTPException(status_code=400, detail="模板表头不正确，请下载最新模板")

    # 数据范围内的服务器，按 IP 建立索引
    scoped_servers = _scope_server_query(db.query(Server), user).all()
    servers_by_ip: dict = {}
    for server in scoped_servers:
        servers_by_ip.setdefault(server.ip, []).append(server)
    server_ids = [s.id for s in scoped_servers]
    account_map = {}
    if server_ids:
        accounts = db.query(ServerSshAccount).filter(
            ServerSshAccount.server_id.in_(server_ids)
        ).all()
        # 用户名按 MySQL 唯一索引的排序规则（大小写不敏感）匹配，
        # 避免仅大小写不同时通过了预校验、却在写库时触发唯一键冲突。
        account_map = {(a.server_id, a.username.casefold()): a for a in accounts}

    records = []
    errors = []
    seen = set()
    default_row_by_server = {}

    for row_index, values in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if row_index - 1 > IMPORT_MAX_ROWS:
            raise HTTPException(status_code=400, detail=f"单次最多导入 {IMPORT_MAX_ROWS} 行")
        if not any(_cell_text(v) for v in values):
            continue

        row = {field: values[index] if index < len(values) else None
               for index, field in field_by_index.items()}
        prefix = f"第{row_index}行"
        ip_raw = _cell_text(row.get("ip"))
        username = _cell_text(row.get("username"))

        if not ip_raw:
            errors.append(f"{prefix}：IP地址不能为空")
            continue
        try:
            ip = str(ipaddress.ip_address(ip_raw))
        except ValueError:
            errors.append(f"{prefix}：IP地址不合法：{ip_raw}")
            continue

        candidates = servers_by_ip.get(ip) or []
        if not candidates:
            errors.append(f"{prefix}：IP {ip} 未纳管，请先在服务器列表中添加")
            continue
        if len(candidates) > 1:
            errors.append(f"{prefix}：IP {ip} 对应多台服务器，无法确定归属")
            continue
        server = candidates[0]

        port_raw = _cell_text(row.get("ssh_port"))
        if port_raw:
            try:
                port = int(float(port_raw))
            except (TypeError, ValueError):
                errors.append(f"{prefix}：SSH端口不合法：{port_raw}")
                continue
            if port != (server.ssh_port or 22):
                errors.append(
                    f"{prefix}：SSH端口 {port} 与服务器登记端口 {server.ssh_port or 22} 不一致"
                )
                continue

        if not username:
            errors.append(f"{prefix}：SSH用户名不能为空")
            continue
        if len(username) > 64:
            errors.append(f"{prefix}：SSH用户名不能超过 64 个字符")
            continue

        key = (server.id, username.casefold())
        if key in seen:
            errors.append(f"{prefix}：账号重复（IP {ip} / {username}）")
            continue
        seen.add(key)

        auth_raw = _cell_text(row.get("auth_type")).lower()
        if auth_raw and auth_raw not in ("password", "passwd", "密码", "key", "密钥", "私钥"):
            errors.append(f"{prefix}：认证方式仅支持 密码 / 密钥")
            continue
        password = _secret_text(row.get("password"))
        private_key = _secret_text(row.get("private_key"))
        if password and private_key:
            errors.append(f"{prefix}：密码与私钥不能同时填写")
            continue

        existing_account = account_map.get(key)
        if existing_account:
            # 大小写不同但指向同一账号时，以库中用户名为准
            username = existing_account.username
        auth_type = "password" if auth_raw in ("password", "passwd", "密码") else (
            "key" if auth_raw in ("key", "密钥", "私钥") else (
                "password" if password else ("key" if private_key else "")
            )
        )
        if not existing_account:
            if not password and not private_key:
                errors.append(f"{prefix}：新增账号必须填写密码或私钥")
                continue
        else:
            if auth_type == "password" and not password and not existing_account.password:
                errors.append(f"{prefix}：该账号未配置密码，认证方式为密码时需填写密码")
                continue
            if auth_type == "key" and not private_key and not existing_account.private_key:
                errors.append(f"{prefix}：该账号未配置私钥，认证方式为密钥时需填写私钥")
                continue

        default_raw = _cell_text(row.get("is_default")).lower()
        if default_raw not in SSH_ACCOUNT_TRUE_VALUES | SSH_ACCOUNT_FALSE_VALUES:
            errors.append(f"{prefix}：是否默认仅支持 是 / 否")
            continue
        is_default = default_raw in SSH_ACCOUNT_TRUE_VALUES
        if is_default:
            if server.id in default_row_by_server:
                errors.append(
                    f"{prefix}：同一服务器只能标记一个默认账号"
                    f"（已标记第 {default_row_by_server[server.id]} 行）"
                )
                continue
            default_row_by_server[server.id] = row_index

        remark = _cell_text(row.get("remark"))
        if len(remark) > 255:
            errors.append(f"{prefix}：备注不能超过 255 个字符")
            continue

        records.append({
            "row": row_index,
            "server": server,
            "username": username,
            "password": password,
            "private_key": private_key,
            "is_default": is_default,
            "remark": remark or None,
            "existing": existing_account,
        })

    # 整批拒绝：任一行校验失败都不写入
    if errors:
        detail = "；".join(errors[:20])
        if len(errors) > 20:
            detail += f"；其余 {len(errors) - 20} 行同类问题未展示"
        raise HTTPException(status_code=400, detail=f"导入校验未通过（{len(errors)} 行）：{detail}")

    added = updated = unchanged = 0
    touched_servers = {}
    items = []
    try:
        for rec in records:
            server = rec["server"]
            account = rec["existing"]
            changed_fields = []
            if account:
                if rec["password"]:
                    account.password = encrypt(rec["password"])
                    account.private_key = None
                    changed_fields.append("密码")
                elif rec["private_key"]:
                    account.private_key = encrypt(rec["private_key"])
                    account.password = None
                    changed_fields.append("私钥")
                if rec["remark"] and rec["remark"] != (account.remark or ""):
                    account.remark = rec["remark"]
                    changed_fields.append("备注")
                if changed_fields:
                    updated += 1
                    status, message = "updated", "已更新：" + "、".join(changed_fields)
                else:
                    unchanged += 1
                    status, message = "unchanged", "无变化"
            else:
                account = ServerSshAccount(
                    server_id=server.id,
                    username=rec["username"],
                    password=encrypt(rec["password"]) if rec["password"] else None,
                    private_key=encrypt(rec["private_key"]) if rec["private_key"] else None,
                    is_default=False,
                    remark=rec["remark"],
                )
                db.add(account)
                db.flush()
                added += 1
                status, message = "added", "新增成功"

            rec["account"] = account
            touched_servers[server.id] = server
            items.append({
                "row": rec["row"],
                "ip": server.ip,
                "username": rec["username"],
                "status": status,
                "message": message,
            })

        # 应用默认账号标记
        for rec in records:
            if rec["is_default"]:
                _set_default_account(db, rec["server"], rec["account"])

        # 保证每台服务器有且仅有一个默认账号，并同步 servers 兼容字段
        for server in touched_servers.values():
            default = _default_account(db, server)
            if not default:
                first = db.query(ServerSshAccount).filter(
                    ServerSshAccount.server_id == server.id
                ).order_by(ServerSshAccount.id).first()
                if first:
                    _set_default_account(db, server, first)
            else:
                _mirror_default_account(server, default)

        db.commit()
    except Exception as e:
        db.rollback()
        raise HTTPException(status_code=400, detail=f"SSH 账号导入失败：{e}")

    log_audit(
        db, user.id, user.username, "IMPORT",
        ip=None, result="success",
        output=f"file={file.filename}, ssh_accounts added={added}, updated={updated}, "
               f"unchanged={unchanged}, rows={len(records)}",
    )

    return ResponseModel(data={
        "total": len(records),
        "added": added,
        "updated": updated,
        "unchanged": unchanged,
        "failed": 0,
        "items": items,
    })
