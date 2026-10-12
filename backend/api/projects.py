from fastapi import APIRouter, Depends, HTTPException, Query, File, UploadFile
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from models.project import Project
from models.subsystem_group_relation import SubsystemGroupRelation
from schemas.project import ProjectCreate, ProjectUpdate, ProjectResponse
from schemas.common import ResponseModel
from services.audit_service import log_audit
from core.database import get_db
from api.dependencies import get_current_user, require_permission
from api.permission_dependencies import assert_project_access, assert_project_manage_capability
from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill, Font
from io import BytesIO
import io

router = APIRouter()


def _project_resource_counts(db: Session, project_ids: list) -> dict:
    """批量统计各项目下的资源数量（子系统数、服务数）。

    返回 {project_id: {"subsystem_count": int, "service_count": int}}。
    一次性聚合查询，避免逐项目 N+1。
    """
    from models.subsystem import Subsystem
    from models.app_service import AppService
    from sqlalchemy import func

    counts = {pid: {"subsystem_count": 0, "service_count": 0} for pid in project_ids}
    if not project_ids:
        return counts

    sub_rows = (
        db.query(Subsystem.project_id, func.count(Subsystem.id))
        .filter(Subsystem.project_id.in_(project_ids))
        .group_by(Subsystem.project_id)
        .all()
    )
    for pid, cnt in sub_rows:
        if pid in counts:
            counts[pid]["subsystem_count"] = cnt

    svc_rows = (
        db.query(AppService.project_id, func.count(AppService.id))
        .filter(AppService.project_id.in_(project_ids))
        .group_by(AppService.project_id)
        .all()
    )
    for pid, cnt in svc_rows:
        if pid in counts:
            counts[pid]["service_count"] = cnt

    return counts


def _project_to_dict(db: Session, project: Project) -> dict:
    """将 Project 对象转换为字典，带组织/部门/负责人名称"""
    from models.organization import Organization
    from models.department import Department
    from models.user import User

    org_name = None
    if project.organization_id:
        org = db.query(Organization).filter(Organization.id == project.organization_id).first()
        if org:
            org_name = org.name

    dept_name = None
    if project.department_id:
        dept = db.query(Department).filter(Department.id == project.department_id).first()
        if dept:
            dept_name = dept.name

    owner_name = None
    if project.owner_id:
        owner = db.query(User).filter(User.id == project.owner_id).first()
        if owner:
            owner_name = owner.job_title or owner.username

    creator_name = None
    if project.creator_id:
        creator = db.query(User).filter(User.id == project.creator_id).first()
        if creator:
            creator_name = creator.job_title or creator.username

    return {
        'id': project.id,
        'name': project.name,
        'project_name': project.name,
        'code': project.code,
        'description': project.description,
        'organization_id': project.organization_id,
        'organization_name': org_name,
        'department_id': project.department_id,
        'department_name': dept_name,
        'creator_id': project.creator_id,
        'creator_name': creator_name,
        'owner_id': project.owner_id,
        'owner_name': owner_name,
        'visibility': project.visibility or 'DEPARTMENT',
    }


@router.get("/projects", response_model=ResponseModel)
async def get_projects(
    page: int = Query(1, ge=1),
    size: int = Query(10, ge=1, le=100),
    name: str = Query(None),
    code: str = Query(None),
    db: Session = Depends(get_db),
    user = Depends(get_current_user)
):
    from services.data_scope_service import DataScopeService

    # 复用 DataScope：非管理员仅能看到有权限的项目
    query = DataScopeService.filter_projects_query(db, user)

    if name:
        query = query.filter(Project.name.like(f"%{name}%"))
    if code:
        query = query.filter(Project.code.like(f"%{code}%"))

    total = query.count()
    projects = query.offset((page - 1) * size).limit(size).all()

    result = [_project_to_dict(db, p) for p in projects]

    # 附加项目下资源数量（子系统数 / 服务数），供项目资源管理页展示
    counts = _project_resource_counts(db, [p.id for p in projects])
    for item in result:
        c = counts.get(item['id'], {})
        item['subsystem_count'] = c.get('subsystem_count', 0)
        item['service_count'] = c.get('service_count', 0)

    return ResponseModel(data={"items": result, "total": total, "page": page, "size": size})

@router.get("/projects/tree", response_model=ResponseModel)
async def get_project_tree(
    db: Session = Depends(get_db),
    user = Depends(get_current_user)
):
    from models.subsystem import Subsystem
    from models.service_group import ServiceGroup
    from sqlalchemy import select, func
    from utils.cache import get_or_set, invalidate_projects_tree, PROJECTS_TREE_CACHE_KEY, PROJECTS_TREE_TTL
    from services.data_scope_service import DataScopeService

    def load_tree_from_db(scoped_user=None):
        """从数据库加载项目树（按 relation.display_order 排序）。

        scoped_user 为 None 表示超管全量；否则按该用户的 DataScope 过滤项目。
        """
        from models.subsystem import Subsystem
        from models.service_group import ServiceGroup
        from models.subsystem_group_relation import SubsystemGroupRelation

        # 1. 查项目（超管全量；普通用户仅可见项目），按 id 排序保证稳定
        if scoped_user is None:
            projects = db.query(Project).order_by(Project.id).all()
        else:
            projects = DataScopeService.filter_projects_query(db, scoped_user).order_by(Project.id).all()

        result = []
        for project in projects:
            project_data = {
                "id": project.id,
                "name": project.name,
                "subsystems": []
            }

            # 2. 查该项目下所有子系统（按 Subsystem.display_order, id 排序）
            subsystems = db.query(Subsystem).filter(
                Subsystem.project_id == project.id
            ).order_by(Subsystem.display_order, Subsystem.id).all()

            for subsystem in subsystems:
                subsystem_data = {
                    "id": subsystem.id,
                    "subsystem_name": subsystem.subsystem_name,
                    "subsystem_code": subsystem.subsystem_code,
                    "service_groups": []
                }

                # 3. 查该子系统下的分组关联（按 relation.display_order 排序）
                relations = db.query(SubsystemGroupRelation).join(
                    ServiceGroup, SubsystemGroupRelation.group_id == ServiceGroup.id
                ).filter(
                    SubsystemGroupRelation.subsystem_id == subsystem.id
                ).order_by(
                    SubsystemGroupRelation.display_order, SubsystemGroupRelation.id
                ).all()

                for relation in relations:
                    if relation.group:
                        subsystem_data["service_groups"].append({
                            "id": relation.group.id,
                            "group_name": relation.group.group_name,
                            "group_code": relation.group.group_code
                        })

                project_data["subsystems"].append(subsystem_data)

            result.append(project_data)
        return result

    # 超管可复用全局缓存（全量项目树，与用户无关）；
    # 非超管直接按当前用户 DataScope 查询，跳过全局缓存，避免跨用户项目泄露
    if DataScopeService._is_super_admin(user):
        tree_data = get_or_set(PROJECTS_TREE_CACHE_KEY, PROJECTS_TREE_TTL, load_tree_from_db)
    else:
        tree_data = load_tree_from_db(user)
    return ResponseModel(data=tree_data)

@router.get("/projects/{project_id}", response_model=ResponseModel)
async def get_project(
    project_id: int,
    db: Session = Depends(get_db),
    user = Depends(get_current_user)
):
    from services.data_scope_service import DataScopeService

    project = db.query(Project).filter(Project.id == project_id).first()
    if not project:
        raise HTTPException(status_code=404, detail="项目不存在")

    # 复用 DataScope：无权限访问该项目时拒绝
    if not DataScopeService.can_access_project(db, user, project_id):
        raise HTTPException(status_code=403, detail="无权访问该项目")

    return ResponseModel(data={
        'id': project.id,
        'name': project.name,
        'project_name': project.name,
        'code': project.code,
        'description': project.description
    })

@router.get("/projects/{project_id}/resources", response_model=ResponseModel)
async def get_project_resources(
    project_id: int,
    db: Session = Depends(get_db),
    user = Depends(get_current_user)
):
    """项目资源归属关系：项目 → 子系统 → 服务。

    复用现有 Project / Subsystem / AppService 结构，不新建模型。
    返回子系统列表（含各自服务）以及未归属子系统的服务。
    """
    from services.data_scope_service import DataScopeService
    from models.subsystem import Subsystem
    from models.app_service import AppService

    project = db.query(Project).filter(Project.id == project_id).first()
    if not project:
        raise HTTPException(status_code=404, detail="项目不存在")

    if not DataScopeService.can_access_project(db, user, project_id):
        raise HTTPException(status_code=403, detail="无权访问该项目")

    subsystems = db.query(Subsystem).filter(
        Subsystem.project_id == project_id
    ).order_by(Subsystem.display_order, Subsystem.id).all()

    services = db.query(AppService).filter(
        AppService.project_id == project_id
    ).order_by(AppService.id).all()

    def _svc_brief(svc: AppService) -> dict:
        return {
            'id': svc.id,
            'func_desc': svc.func_desc,
            'module': svc.module,
            'ip': svc.ip,
            'port': svc.port,
            'service_type': svc.service_type,
            'deploy_type': svc.deploy_type,
            'subsystem_id': svc.subsystem_id,
            'group_id': svc.group_id,
            'owner': svc.owner,
        }

    # 按子系统归组服务
    svc_by_subsystem = {}
    unassigned = []
    for svc in services:
        if svc.subsystem_id:
            svc_by_subsystem.setdefault(svc.subsystem_id, []).append(_svc_brief(svc))
        else:
            unassigned.append(_svc_brief(svc))

    subsystem_list = []
    for sub in subsystems:
        subsystem_list.append({
            'id': sub.id,
            'subsystem_name': sub.subsystem_name,
            'subsystem_code': sub.subsystem_code,
            'description': sub.description,
            'display_order': sub.display_order,
            'service_count': len(svc_by_subsystem.get(sub.id, [])),
            'services': svc_by_subsystem.get(sub.id, []),
        })

    return ResponseModel(data={
        'project_id': project.id,
        'project_name': project.name,
        'subsystem_count': len(subsystems),
        'service_count': len(services),
        'subsystems': subsystem_list,
        'unassigned_services': unassigned,
    })

@router.get("/projects/{project_id}/subsystems", response_model=ResponseModel)
async def get_project_subsystems(
    project_id: int,
    db: Session = Depends(get_db),
    user = Depends(get_current_user)
):
    """获取项目子系统列表"""
    from models.subsystem import Subsystem

    project = db.query(Project).filter(Project.id == project_id).first()
    if not project:
        raise HTTPException(status_code=404, detail="项目不存在")
    assert_project_access(db, user, project_id)

    subsystems = db.query(Subsystem).filter(Subsystem.project_id == project_id).all()
    result = []
    for subsystem in subsystems:
        result.append({
            'id': subsystem.id,
            'subsystem_name': subsystem.subsystem_name,
            'subsystem_code': subsystem.subsystem_code,
            'description': subsystem.description
        })
    
    return ResponseModel(data=result)

@router.post("/projects", response_model=ResponseModel)
async def create_project(
    project_create: ProjectCreate,
    db: Session = Depends(get_db),
    user = Depends(require_permission("project_manage"))
):
    existing = db.query(Project).filter(Project.name == project_create.name).first()
    if existing:
        raise HTTPException(status_code=400, detail="项目名称已存在")

    # 校验所属组织存在性
    organization_id = project_create.organization_id
    if not organization_id:
        # 未指定组织时，默认使用当前用户所在组织
        organization_id = getattr(user, "organization_id", None)
    if organization_id:
        from models.organization import Organization
        org = db.query(Organization).filter(Organization.id == organization_id).first()
        if not org:
            raise HTTPException(status_code=400, detail="所属组织不存在")

    # 校验部门归属与组织一致
    if project_create.department_id and organization_id:
        from models.department import Department
        dept = db.query(Department).filter(Department.id == project_create.department_id).first()
        if not dept:
            raise HTTPException(status_code=400, detail="所属部门不存在")
        if dept.organization_id != organization_id:
            raise HTTPException(status_code=400, detail="所属部门不属于该组织")

    new_project = Project(
        name=project_create.name,
        code=project_create.code,
        description=project_create.description,
        organization_id=organization_id,
        department_id=project_create.department_id,
        visibility=project_create.visibility or "DEPARTMENT",
        owner_id=project_create.owner_id,
        creator_id=user.id,
    )

    db.add(new_project)
    db.commit()
    db.refresh(new_project)

    # 清除项目树缓存
    from utils.cache import invalidate_projects_tree
    invalidate_projects_tree()

    log_audit(db, user.id, user.username, "PROJECT_CREATE",
              result="success", output=f"创建项目: {project_create.name}")

    return ResponseModel(data=_project_to_dict(db, new_project))

@router.put("/projects/{project_id}", response_model=ResponseModel)
async def update_project(
    project_id: int,
    project_update: ProjectUpdate,
    db: Session = Depends(get_db),
    user = Depends(require_permission("project_manage"))
):
    project = db.query(Project).filter(Project.id == project_id).first()
    if not project:
        raise HTTPException(status_code=404, detail="项目不存在")

    # 数据范围校验：仅有 project_manage 不足以修改无权限项目
    assert_project_access(db, user, project_id)
    # 管理能力校验：仅通过项目共享/公开可见获得访问的用户不得修改项目属性与可见范围
    assert_project_manage_capability(db, user, project_id)

    if project_update.name and project_update.name != project.name:
        existing = db.query(Project).filter(Project.name == project_update.name).first()
        if existing:
            raise HTTPException(status_code=400, detail="项目名称已存在")
        project.name = project_update.name

    if project_update.code is not None:
        project.code = project_update.code
    if project_update.description is not None:
        project.description = project_update.description

    # 组织可用于首次设置；已设置组织时不允许修改（由前端锁定控制）
    if project_update.organization_id is not None:
        if project.organization_id and project.organization_id != project_update.organization_id:
            raise HTTPException(status_code=400, detail="项目已关联组织，不可修改")
        if project_update.organization_id:
            from models.organization import Organization
            org = db.query(Organization).filter(Organization.id == project_update.organization_id).first()
            if not org:
                raise HTTPException(status_code=400, detail="所属组织不存在")
        project.organization_id = project_update.organization_id

    # 部门可修改，但必须属于项目所在组织
    if project_update.department_id is not None:
        if project_update.department_id:
            from models.department import Department
            dept = db.query(Department).filter(Department.id == project_update.department_id).first()
            if not dept:
                raise HTTPException(status_code=400, detail="所属部门不存在")
            effective_org_id = project.organization_id
            if effective_org_id and dept.organization_id != effective_org_id:
                raise HTTPException(status_code=400, detail="所属部门不属于该项目所在组织")
        project.department_id = project_update.department_id

    if project_update.visibility is not None:
        project.visibility = project_update.visibility
    if project_update.owner_id is not None:
        project.owner_id = project_update.owner_id

    db.commit()
    db.refresh(project)

    # 清除项目树缓存
    from utils.cache import invalidate_projects_tree
    invalidate_projects_tree()

    log_audit(db, user.id, user.username, "PROJECT_UPDATE",
              result="success", output=f"更新项目: {project.name}")

    return ResponseModel(data=_project_to_dict(db, project))

@router.delete("/projects/{project_id}", response_model=ResponseModel)
async def delete_project(
    project_id: int,
    db: Session = Depends(get_db),
    user = Depends(require_permission("project_manage"))
):
    project = db.query(Project).filter(Project.id == project_id).first()
    if not project:
        raise HTTPException(status_code=404, detail="项目不存在")

    # 数据范围校验：仅有 project_manage 不足以删除无权限项目
    assert_project_access(db, user, project_id)
    # 管理能力校验：仅通过项目共享/公开可见获得访问的用户不得删除项目
    assert_project_manage_capability(db, user, project_id)

    project_name = project.name

    # 1. 先清理没有数据库外键级联的关联数据
    # AppService 没有数据库级外键约束，需要手动清理
    from models.app_service import AppService
    db.query(AppService).filter(AppService.project_id == project_id).delete(synchronize_session=False)

    # 2. ProjectMember 有 CASCADE 外键，但显式删除更安全
    from models.project_member import ProjectMember
    db.query(ProjectMember).filter(ProjectMember.project_id == project_id).delete(synchronize_session=False)

    # 3. ProjectPermission 有 CASCADE 外键
    from models.project_permission import ProjectPermission
    db.query(ProjectPermission).filter(ProjectPermission.project_id == project_id).delete(synchronize_session=False)

    # 4. RoleProjectPermission 无数据库级外键
    from models.role_project_permission import RoleProjectPermission
    db.query(RoleProjectPermission).filter(RoleProjectPermission.project_id == project_id).delete(synchronize_session=False)

    # 5. ResourceGroup 有 CASCADE 外键
    from models.resource_group import ResourceGroup
    db.query(ResourceGroup).filter(ResourceGroup.project_id == project_id).delete(synchronize_session=False)

    # 6. 清理 SubsystemGroupRelation（Subsystem 的级联）
    from models.subsystem import Subsystem
    from models.subsystem_group_relation import SubsystemGroupRelation
    subsystem_ids = [s[0] for s in db.query(Subsystem.id).filter(Subsystem.project_id == project_id).all()]
    if subsystem_ids:
        db.query(SubsystemGroupRelation).filter(SubsystemGroupRelation.subsystem_id.in_(subsystem_ids)).delete(synchronize_session=False)

    # 7. 最后删除项目（CASCADE 外键会自动处理 Subsystem）
    db.delete(project)
    db.commit()

    # 清除项目树缓存
    from utils.cache import invalidate_projects_tree
    invalidate_projects_tree()

    log_audit(db, user.id, user.username, "PROJECT_DELETE",
              result="success", output=f"删除项目: {project_name}")

    return ResponseModel(data={"id": project_id})

@router.get("/projects/template")
async def download_project_template(
    user = Depends(require_permission("project_manage"))
):
    """下载项目导入模板"""
    headers = ['项目名称', '项目编码', '项目描述']

    wb = Workbook()
    ws = wb.active
    ws.title = "项目导入模板"

    header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
    header_font = Font(color='FFFFFF', bold=True)

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=projects_template.xlsx"}
    )


@router.get("/projects/export")
async def export_projects(
    db: Session = Depends(get_db),
    user = Depends(require_permission("project_manage"))
):
    """导出项目列表"""
    from services.data_scope_service import DataScopeService

    # 与项目列表使用一致的数据范围，避免导出越权项目
    projects = DataScopeService.filter_projects_query(db, user).all()

    headers = ['name', 'code', 'description']

    wb = Workbook()
    ws = wb.active
    ws.title = "项目列表"

    header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
    header_font = Font(color='FFFFFF', bold=True)

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font

    for row_idx, project in enumerate(projects, 2):
        ws.cell(row=row_idx, column=1, value=project.name)
        ws.cell(row=row_idx, column=2, value=project.code or '')
        ws.cell(row=row_idx, column=3, value=project.description or '')

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=projects.xlsx"}
    )

@router.post("/projects/import", response_model=ResponseModel)
async def import_projects(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    user = Depends(require_permission("project_manage"))
):
    if not file.filename.endswith('.xlsx'):
        raise HTTPException(status_code=400, detail="仅支持xlsx格式文件")

    try:
        file_content = await file.read()
        wb = load_workbook(filename=io.BytesIO(file_content), read_only=True)
        ws = wb.active

        headers = []
        for col in range(1, ws.max_column + 1):
            header = ws.cell(row=1, column=col).value
            headers.append(header.strip() if header else '')

        rows = []
        for row in range(2, ws.max_row + 1):
            row_data = {}
            for col in range(1, ws.max_column + 1):
                header = headers[col - 1]
                value = ws.cell(row=row, column=col).value
                if value is not None:
                    row_data[header] = str(value) if isinstance(value, (int, float)) else value
            if row_data:
                rows.append(row_data)
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"文件解析失败: {str(e)}")

    added = 0
    updated = 0
    failed = 0
    failed_items = []

    for idx, row in enumerate(rows):
        try:
            name = str(row.get('name', '')).strip()
            if not name:
                failed += 1
                failed_items.append({"row": idx + 2, "reason": "缺少项目名称"})
                continue

            existing = db.query(Project).filter(Project.name == name).first()

            if existing:
                if 'code' in row:
                    existing.code = str(row.get('code', '')).strip() or None
                if 'description' in row:
                    existing.description = str(row.get('description', '')).strip() or None
                db.commit()
                updated += 1
            else:
                new_project = Project(
                    name=name,
                    code=str(row.get('code', '')).strip() or None,
                    description=str(row.get('description', '')).strip() or None
                )
                db.add(new_project)
                db.commit()
                added += 1
        except Exception as e:
            failed += 1
            failed_items.append({"row": idx + 2, "reason": str(e)})

    log_audit(db, user.id, user.username, "PROJECT_IMPORT",
              result="success",
              output=f"导入完成: 新增{added}条, 更新{updated}条, 失败{failed}条")

    return ResponseModel(data={
        "success": True,
        "message": f"导入完成: 新增{added}条, 更新{updated}条, 失败{failed}条",
        "added": added,
        "updated": updated,
        "failed": failed,
        "failed_items": failed_items
    })

@router.get("/projects/list", response_model=ResponseModel)
async def list_projects(
    db: Session = Depends(get_db),
    user = Depends(get_current_user)
):
    from services.data_scope_service import DataScopeService

    # 复用 DataScope：非管理员仅能看到有权限的项目
    projects = DataScopeService.filter_projects_query(db, user).all()
    result = [{"id": p.id, "name": p.name, "project_name": p.name} for p in projects]
    return ResponseModel(data=result)
