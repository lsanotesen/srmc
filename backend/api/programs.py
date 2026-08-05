from fastapi import APIRouter, Depends, HTTPException, Query, File, UploadFile
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from models.program import Program
from schemas.program import (
    ProgramCreate, ProgramUpdate, ProgramResponse, 
    ProgramOperationResult, ProgramImportResult
)
from schemas.common import ResponseModel
from utils.crypto import encrypt, decrypt
from utils.program_status import get_program_status, get_batch_program_status
from openpyxl import load_workbook, Workbook
from openpyxl.styles import PatternFill, Font
from services.audit_service import log_audit
from core.database import get_db
from api.dependencies import get_current_user, require_permission
from utils.ssh_pool import ssh_pool
from typing import List, Optional
import asyncio
import io

def format_script_command(script: str) -> str:
    """格式化脚本命令，支持多行脚本，将换行转换为 ; 连接（确保所有命令都执行）"""
    lines = script.strip().split('\n')
    formatted_lines = []
    for line in lines:
        line = line.strip()
        if line:
            # 如果是脚本文件名（不含空格）且没有路径前缀，添加 ./ 前缀
            if ' ' not in line and not line.startswith('./') and not line.startswith('/') and not line.startswith('~'):
                line = f"./{line}"
            formatted_lines.append(line)
    return '; '.join(formatted_lines)

router = APIRouter()

@router.get("/programs", response_model=ResponseModel)
async def get_programs(
    func_desc: Optional[str] = None,
    module: Optional[str] = None,
    ip: Optional[str] = None,
    owner: Optional[str] = None,
    page: int = Query(1, ge=1),
    size: int = Query(10, ge=1, le=100),
    db: Session = Depends(get_db),
    user = Depends(get_current_user)
):
    query = db.query(Program)
    
    if func_desc:
        query = query.filter(Program.func_desc.like(f"%{func_desc}%"))
    if module:
        query = query.filter(Program.module.like(f"%{module}%"))
    if ip:
        query = query.filter(Program.ip.like(f"%{ip}%"))
    if owner:
        query = query.filter(Program.owner.like(f"%{owner}%"))
    
    total = query.count()
    programs = query.offset((page - 1) * size).limit(size).all()
    
    result = []
    for program in programs:
        program_dict = {
            'id': program.id,
            'func_desc': program.func_desc,
            'module': program.module,
            'ip': program.ip,
            'username': program.username,
            'program_path': program.program_path,
            'start_script': program.start_script,
            'stop_script': program.stop_script,
            'restart_script': program.restart_script,
            'log_path': program.log_path,
            'port': program.port,
            'owner': program.owner,
            'remark': program.remark,
            'status': None,
            'created_at': program.created_at,
            'updated_at': program.updated_at
        }
        result.append(program_dict)
    
    return ResponseModel(data={"items": result, "total": total, "page": page, "size": size})

@router.get("/programs/{program_id}", response_model=ResponseModel)
async def get_program(
    program_id: int,
    db: Session = Depends(get_db),
    user = Depends(get_current_user)
):
    program = db.query(Program).filter(Program.id == program_id).first()
    if not program:
        raise HTTPException(status_code=404, detail="程序不存在")
    
    return ResponseModel(data=ProgramResponse.model_validate(program))

@router.post("/programs", response_model=ResponseModel)
async def create_program(
    program: ProgramCreate,
    db: Session = Depends(get_db),
    user = Depends(require_permission("program:manage"))
):
    existing = db.query(Program).filter(
        (Program.func_desc == program.func_desc) &
        (Program.ip == str(program.ip)) &
        (Program.program_path == program.program_path)
    ).first()
    if existing:
        raise HTTPException(status_code=400, detail="程序已存在（功能描述+IP+程序路径重复）")
    
    new_program = Program(
        func_desc=program.func_desc,
        module=program.module,
        ip=str(program.ip),
        username=program.username,
        password=encrypt(program.password),
        program_path=program.program_path,
        start_script=program.start_script,
        stop_script=program.stop_script,
        restart_script=program.restart_script,
        log_path=program.log_path,
        port=program.port,
        owner=program.owner,
        remark=program.remark
    )
    
    db.add(new_program)
    db.commit()
    db.refresh(new_program)
    
    log_audit(db, user.id, user.username, "PROGRAM_CREATE", 
              result="success", output=f"创建程序: {program.func_desc}")
    
    return ResponseModel(data=ProgramResponse.model_validate(new_program))

@router.put("/programs/{program_id}", response_model=ResponseModel)
async def update_program(
    program_id: int,
    program: ProgramUpdate,
    db: Session = Depends(get_db),
    user = Depends(require_permission("program:manage"))
):
    existing = db.query(Program).filter(Program.id == program_id).first()
    if not existing:
        raise HTTPException(status_code=404, detail="程序不存在")
    
    if program.func_desc is not None:
        existing.func_desc = program.func_desc
    if program.module is not None:
        existing.module = program.module
    if program.ip is not None:
        existing.ip = str(program.ip)
    if program.username is not None:
        existing.username = program.username
    if program.password is not None and program.password != '':
        existing.password = encrypt(program.password)
    if program.program_path is not None:
        existing.program_path = program.program_path
    if program.start_script is not None:
        existing.start_script = program.start_script
    if program.stop_script is not None:
        existing.stop_script = program.stop_script
    if program.restart_script is not None:
        existing.restart_script = program.restart_script
    if program.log_path is not None:
        existing.log_path = program.log_path
    if program.port is not None:
        existing.port = program.port
    if program.owner is not None:
        existing.owner = program.owner
    if program.remark is not None:
        existing.remark = program.remark
    
    db.commit()
    db.refresh(existing)
    
    log_audit(db, user.id, user.username, "PROGRAM_UPDATE", 
              result="success", output=f"更新程序: {existing.func_desc}")
    
    return ResponseModel(data=ProgramResponse.model_validate(existing))

@router.delete("/programs/{program_id}", response_model=ResponseModel)
async def delete_program(
    program_id: int,
    db: Session = Depends(get_db),
    user = Depends(require_permission("program:manage"))
):
    program = db.query(Program).filter(Program.id == program_id).first()
    if not program:
        raise HTTPException(status_code=404, detail="程序不存在")
    
    func_desc = program.func_desc
    db.delete(program)
    db.commit()
    
    log_audit(db, user.id, user.username, "PROGRAM_DELETE", 
              result="success", output=f"删除程序: {func_desc}")
    
    return ResponseModel(message="删除成功")

@router.post("/programs/status/batch", response_model=ResponseModel)
async def batch_get_status(
    program_ids: List[int],
    db: Session = Depends(get_db),
    user = Depends(get_current_user)
):
    programs = db.query(Program).filter(Program.id.in_(program_ids)).all()
    
    # 准备程序信息（包含解密后的密码）
    progs_info = []
    for prog in programs:
        progs_info.append({
            'id': prog.id,
            'ip': prog.ip,
            'username': prog.username,
            'password': decrypt(prog.password),
            'port': prog.port,
            'program_path': prog.program_path
        })
    
    status_dict = await get_batch_program_status(progs_info)
    return ResponseModel(data=status_dict)

async def execute_program_script(
    program: Program, 
    script_type: str,
    db: Session,
    user
) -> ProgramOperationResult:
    """执行程序脚本（启动/停止/重启）"""
    try:
        password = decrypt(program.password)
        conn = await ssh_pool.get_connection(program.ip, 22, program.username, password)
        
        if not conn:
            log_audit(db, user.id, user.username, f"PROGRAM_{script_type.upper()}",
                      result="failed", output=f"SSH连接失败: {program.ip}")
            return ProgramOperationResult(success=False, message="SSH连接失败")
        
        # 构建命令
        if script_type == "start":
            script = program.start_script
            action = "启动"
        elif script_type == "stop":
            script = program.stop_script
            action = "停止"
        elif script_type == "restart":
            script = program.restart_script
            action = "重启"
        else:
            return ProgramOperationResult(success=False, message="未知操作类型")
        
        if not script:
            if script_type == "restart":
                # 没有重启脚本，先停止再启动
                stop_script = program.stop_script
                start_script = program.start_script
                if not stop_script or not start_script:
                    return ProgramOperationResult(success=False, message="缺少停止或启动脚本")
                
                # 执行停止
                stop_cmd = f"cd {program.program_path} && {format_script_command(stop_script)}"
                stop_output, stop_error, stop_success = await conn.execute_command(stop_cmd)
                
                if not stop_success:
                    log_audit(db, user.id, user.username, "PROGRAM_STOP",
                              result="failed", output=f"停止失败: {stop_error}")
                    return ProgramOperationResult(
                        success=False, 
                        message="停止失败",
                        output=stop_error
                    )
                
                await asyncio.sleep(2)
                
                # 执行启动
                start_cmd = f"cd {program.program_path} && {format_script_command(start_script)}"
                output, error, success = await conn.execute_command(start_cmd)
            else:
                return ProgramOperationResult(success=False, message=f"缺少{action}脚本")
        else:
            command = f"cd {program.program_path} && {format_script_command(script)}"
            output, error, success = await conn.execute_command(command)
        
        if success:
            log_audit(db, user.id, user.username, f"PROGRAM_{script_type.upper()}",
                      result="success", output=f"{action}成功: {program.func_desc}")
            
            # 等待2秒后检测状态
            await asyncio.sleep(2)
            status = await get_program_status(
                program.ip, program.username, password, program.port, program.program_path
            )
            
            return ProgramOperationResult(
                success=True,
                message=f"{action}成功，当前状态: {status}",
                output=output[:500] if output else None
            )
        else:
            log_audit(db, user.id, user.username, f"PROGRAM_{script_type.upper()}",
                      result="failed", output=f"{action}失败: {error}")
            return ProgramOperationResult(
                success=False,
                message=f"{action}失败",
                output=error[:500] if error else None
            )
    except Exception as e:
        log_audit(db, user.id, user.username, f"PROGRAM_{script_type.upper()}",
                  result="failed", output=f"{action}异常: {str(e)}")
        return ProgramOperationResult(success=False, message=f"{action}异常: {str(e)}")

@router.post("/programs/{program_id}/start", response_model=ResponseModel)
async def start_program(
    program_id: int,
    db: Session = Depends(get_db),
    user = Depends(require_permission("program:operate"))
):
    program = db.query(Program).filter(Program.id == program_id).first()
    if not program:
        raise HTTPException(status_code=404, detail="程序不存在")
    
    result = await execute_program_script(program, "start", db, user)
    return ResponseModel(data=result.dict())

@router.post("/programs/{program_id}/stop", response_model=ResponseModel)
async def stop_program(
    program_id: int,
    db: Session = Depends(get_db),
    user = Depends(require_permission("program:operate"))
):
    program = db.query(Program).filter(Program.id == program_id).first()
    if not program:
        raise HTTPException(status_code=404, detail="程序不存在")
    
    result = await execute_program_script(program, "stop", db, user)
    return ResponseModel(data=result.dict())

@router.post("/programs/{program_id}/restart", response_model=ResponseModel)
async def restart_program(
    program_id: int,
    db: Session = Depends(get_db),
    user = Depends(require_permission("program:operate"))
):
    program = db.query(Program).filter(Program.id == program_id).first()
    if not program:
        raise HTTPException(status_code=404, detail="程序不存在")
    
    result = await execute_program_script(program, "restart", db, user)
    return ResponseModel(data=result.dict())

@router.get("/programs/{program_id}/log", response_model=ResponseModel)
async def get_program_log(
    program_id: int,
    lines: int = Query(200, ge=10, le=2000),
    db: Session = Depends(get_db),
    user = Depends(get_current_user)
):
    program = db.query(Program).filter(Program.id == program_id).first()
    if not program:
        raise HTTPException(status_code=404, detail="程序不存在")
    
    if not program.log_path:
        return ResponseModel(data={"content": "", "log_path": program.log_path})
    
    try:
        password = decrypt(program.password)
        conn = await ssh_pool.get_connection(program.ip, 22, program.username, password)
        
        if not conn:
            return ResponseModel(data={"content": "SSH连接失败", "log_path": program.log_path})
        
        # 判断log_path是文件还是目录
        check_cmd = f"if [ -d '{program.log_path}' ]; then echo 'dir'; elif [ -f '{program.log_path}' ]; then echo 'file'; else echo 'notfound'; fi"
        output, error, success = await conn.execute_command(check_cmd)
        
        if not success:
            return ResponseModel(data={"content": f"检查路径失败: {error}", "log_path": program.log_path})
        
        path_type = output.strip()
        
        if path_type == "dir":
            # 目录：获取最新的.log文件
            cmd = f"ls -t {program.log_path}/*.log 2>/dev/null | head -1"
            output, error, success = await conn.execute_command(cmd)
            if not success or not output.strip():
                return ResponseModel(data={"content": "未找到日志文件", "log_path": program.log_path})
            log_file = output.strip()
        elif path_type == "file":
            log_file = program.log_path
        else:
            return ResponseModel(data={"content": "日志路径不存在", "log_path": program.log_path})
        
        # 读取日志
        cmd = f"tail -n {lines} {log_file}"
        output, error, success = await conn.execute_command(cmd)
        
        if success:
            return ResponseModel(data={"content": output, "log_path": log_file})
        else:
            return ResponseModel(data={"content": f"读取日志失败: {error}", "log_path": log_file})
    except Exception as e:
        return ResponseModel(data={"content": f"获取日志异常: {str(e)}", "log_path": program.log_path})

@router.get("/programs/{program_id}/log/download")
async def download_program_log(
    program_id: int,
    db: Session = Depends(get_db),
    user = Depends(get_current_user)
):
    program = db.query(Program).filter(Program.id == program_id).first()
    if not program:
        raise HTTPException(status_code=404, detail="程序不存在")
    
    if not program.log_path:
        raise HTTPException(status_code=400, detail="未配置日志路径")
    
    try:
        password = decrypt(program.password)
        conn = await ssh_pool.get_connection(program.ip, 22, program.username, password)
        
        if not conn:
            raise HTTPException(status_code=500, detail="SSH连接失败")
        
        # 判断log_path是文件还是目录
        check_cmd = f"if [ -d '{program.log_path}' ]; then echo 'dir'; elif [ -f '{program.log_path}' ]; then echo 'file'; else echo 'notfound'; fi"
        output, error, success = await conn.execute_command(check_cmd)
        
        if not success:
            raise HTTPException(status_code=500, detail=f"检查路径失败: {error}")
        
        path_type = output.strip()
        
        if path_type == "dir":
            cmd = f"ls -t {program.log_path}/*.log 2>/dev/null | head -1"
            output, error, success = await conn.execute_command(cmd)
            if not success or not output.strip():
                raise HTTPException(status_code=400, detail="未找到日志文件")
            log_file = output.strip()
        elif path_type == "file":
            log_file = program.log_path
        else:
            raise HTTPException(status_code=400, detail="日志路径不存在")
        
        # 读取完整日志
        cmd = f"cat {log_file}"
        output, error, success = await conn.execute_command(cmd)
        
        if not success:
            raise HTTPException(status_code=500, detail=f"读取日志失败: {error}")
        
        filename = f"{program.func_desc}.log"
        return StreamingResponse(
            io.StringIO(output),
            media_type="text/plain",
            headers={"Content-Disposition": f"attachment; filename={filename}"}
        )
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"下载日志失败: {str(e)}")

@router.get("/programs/export")
async def export_programs(
    db: Session = Depends(get_db),
    user = Depends(require_permission("program:manage"))
):
    programs = db.query(Program).all()
    
    headers = ['func_desc', 'module', 'ip', 'username', 'password', 'program_path', 
               'start_script', 'stop_script', 'restart_script', 'log_path', 'port', 'owner', 'remark']
    
    wb = Workbook()
    ws = wb.active
    ws.title = "程序列表"
    
    header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
    header_font = Font(color='FFFFFF', bold=True)
    
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
    
    for row_idx, program in enumerate(programs, 2):
        ws.cell(row=row_idx, column=1, value=program.func_desc)
        ws.cell(row=row_idx, column=2, value=program.module or '')
        ws.cell(row=row_idx, column=3, value=program.ip)
        ws.cell(row=row_idx, column=4, value=program.username)
        ws.cell(row=row_idx, column=5, value='')  # 不导出密码
        ws.cell(row=row_idx, column=6, value=program.program_path)
        ws.cell(row=row_idx, column=7, value=program.start_script or '')
        ws.cell(row=row_idx, column=8, value=program.stop_script or '')
        ws.cell(row=row_idx, column=9, value=program.restart_script or '')
        ws.cell(row=row_idx, column=10, value=program.log_path or '')
        ws.cell(row=row_idx, column=11, value=program.port or '')
        ws.cell(row=row_idx, column=12, value=program.owner or '')
        ws.cell(row=row_idx, column=13, value=program.remark or '')
    
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=programs.xlsx"}
    )

@router.post("/programs/import", response_model=ResponseModel)
async def import_programs(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    user = Depends(require_permission("program:manage"))
):
    if not file.filename.endswith('.xlsx'):
        raise HTTPException(status_code=400, detail="仅支持xlsx格式文件")
    
    try:
        file_content = await file.read()
        wb = load_workbook(filename=io.BytesIO(file_content), read_only=True)
        ws = wb.active
        
        headers = []
        for col in range(1, ws.max_column + 1):
            header = ws.cell(row=1, column=col).value
            headers.append(header.strip() if header else '')
        
        rows = []
        for row in range(2, ws.max_row + 1):
            row_data = {}
            for col in range(1, ws.max_column + 1):
                header = headers[col - 1]
                value = ws.cell(row=row, column=col).value
                if value is not None:
                    row_data[header] = str(value) if isinstance(value, (int, float)) else value
            if row_data:
                rows.append(row_data)
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"文件解析失败: {str(e)}")
    
    added = 0
    updated = 0
    failed = 0
    failed_items = []
    
    for idx, row in enumerate(rows):
        try:
            func_desc = str(row.get('func_desc', '')).strip()
            ip = str(row.get('ip', '')).strip()
            program_path = str(row.get('program_path', '')).strip()
            
            if not func_desc or not ip or not program_path:
                failed += 1
                failed_items.append({"row": idx + 2, "reason": "缺少必填字段"})
                continue
            
            # 检查是否已存在
            existing = db.query(Program).filter(
                (Program.func_desc == func_desc) &
                (Program.ip == ip) &
                (Program.program_path == program_path)
            ).first()
            
            password = str(row.get('password', '')).strip()
            
            if existing:
                # 更新现有记录
                if 'module' in row:
                    existing.module = str(row.get('module', '')).strip() or None
                if 'username' in row:
                    existing.username = str(row.get('username', '')).strip()
                if password:
                    existing.password = encrypt(password)
                if 'start_script' in row:
                    existing.start_script = str(row.get('start_script', '')).strip() or None
                if 'stop_script' in row:
                    existing.stop_script = str(row.get('stop_script', '')).strip() or None
                if 'restart_script' in row:
                    existing.restart_script = str(row.get('restart_script', '')).strip() or None
                if 'log_path' in row:
                    existing.log_path = str(row.get('log_path', '')).strip() or None
                if 'port' in row and str(row.get('port', '')).strip():
                    existing.port = int(row.get('port'))
                if 'owner' in row:
                    existing.owner = str(row.get('owner', '')).strip() or None
                if 'remark' in row:
                    existing.remark = str(row.get('remark', '')).strip() or None
                
                db.commit()
                updated += 1
            else:
                # 创建新记录
                new_program = Program(
                    func_desc=func_desc,
                    module=str(row.get('module', '')).strip() or None,
                    ip=ip,
                    username=str(row.get('username', '')).strip() or 'root',
                    password=encrypt(password) if password else encrypt('password'),
                    program_path=program_path,
                    start_script=str(row.get('start_script', '')).strip() or None,
                    stop_script=str(row.get('stop_script', '')).strip() or None,
                    restart_script=str(row.get('restart_script', '')).strip() or None,
                    log_path=str(row.get('log_path', '')).strip() or None,
                    port=int(row.get('port')) if row.get('port') and str(row.get('port', '')).strip() else None,
                    owner=str(row.get('owner', '')).strip() or None,
                    remark=str(row.get('remark', '')).strip() or None
                )
                
                db.add(new_program)
                db.commit()
                added += 1
        except Exception as e:
            failed += 1
            failed_items.append({"row": idx + 2, "reason": str(e)})
    
    log_audit(db, user.id, user.username, "PROGRAM_IMPORT",
              result="success", 
              output=f"导入完成: 新增{added}条, 更新{updated}条, 失败{failed}条")
    
    return ResponseModel(data={
        "success": True,
        "message": f"导入完成: 新增{added}条, 更新{updated}条, 失败{failed}条",
        "added": added,
        "updated": updated,
        "failed": failed,
        "failed_items": failed_items
    })