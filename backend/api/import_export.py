from fastapi import APIRouter, Depends, File, UploadFile
from sqlalchemy.orm import Session
from models.service import Service
from schemas.common import ResponseModel
from services.import_export_service import import_services
from utils.excel_utils import generate_excel_template
from core.database import get_db
from api.dependencies import get_current_user, require_permission

router = APIRouter()

@router.post("/import/services", response_model=ResponseModel)
async def import_services_endpoint(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    user = Depends(require_permission("import"))
):
    if not file.filename.endswith('.xlsx'):
        return ResponseModel(code=1, message="Only Excel files (.xlsx) are supported")
    
    content = await file.read()
    result = await import_services(content, db, user.id, user.username)
    
    return ResponseModel(data=result.dict())

@router.get("/import/template")
async def download_template(user = Depends(require_permission("import"))):
    from fastapi.responses import StreamingResponse
    output = generate_excel_template()
    return StreamingResponse(output, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                            headers={"Content-Disposition": "attachment; filename=service_import_template.xlsx"})

@router.get("/export/services", response_model=ResponseModel)
async def export_services(db: Session = Depends(get_db), user = Depends(require_permission("export"))):
    from fastapi.responses import StreamingResponse
    from openpyxl import Workbook
    from io import BytesIO
    
    services = db.query(Service).all()
    
    wb = Workbook()
    ws = wb.active
    ws.title = "Services"
    
    headers = [
        'service_name', 'service_code', 'service_type', 'module', 'environment',
        'ip', 'port', 'service_path', 'work_dir', 'start_script',
        'stop_script', 'restart_script', 'log_path', 'check_type',
        'check_keyword', 'pid_file', 'owner', 'remark'
    ]
    
    for col, header in enumerate(headers, 1):
        ws.cell(row=1, column=col, value=header)
    
    for row_idx, service in enumerate(services, 2):
        ws.cell(row=row_idx, column=1, value=service.service_name)
        ws.cell(row=row_idx, column=2, value=service.service_code)
        ws.cell(row=row_idx, column=3, value=service.service_type)
        ws.cell(row=row_idx, column=4, value=service.module or '')
        ws.cell(row=row_idx, column=5, value=service.environment)
        ws.cell(row=row_idx, column=6, value=service.ip)
        ws.cell(row=row_idx, column=7, value=service.port or '')
        ws.cell(row=row_idx, column=8, value=service.service_path or '')
        ws.cell(row=row_idx, column=9, value=service.work_dir or '')
        ws.cell(row=row_idx, column=10, value=service.start_script or '')
        ws.cell(row=row_idx, column=11, value=service.stop_script or '')
        ws.cell(row=row_idx, column=12, value=service.restart_script or '')
        ws.cell(row=row_idx, column=13, value=service.log_path or '')
        ws.cell(row=row_idx, column=14, value=service.check_type or '')
        ws.cell(row=row_idx, column=15, value=service.check_keyword or '')
        ws.cell(row=row_idx, column=16, value=service.pid_file or '')
        ws.cell(row=row_idx, column=17, value=service.owner or '')
        ws.cell(row=row_idx, column=18, value=service.remark or '')
    
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    
    return StreamingResponse(output, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                            headers={"Content-Disposition": "attachment; filename=services_export.xlsx"})
